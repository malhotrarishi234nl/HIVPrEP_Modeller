"""Testable service boundary used by the Streamlit workbench UI.

This module can be tested without installing Streamlit.  The optional workbook
preview uses openpyxl.  Simulation dispatch is delegated to
``model_workbench.service.WorkbenchService`` when a compatible recipe is
selected.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import csv
import hashlib
import importlib.metadata
import json
import math
from pathlib import Path
import re
import subprocess
import sys
from typing import Any, Mapping, Sequence
import zipfile
from xml.etree import ElementTree

from model_workbench.workers import (
    ALL_AVAILABLE_WORKERS,
    MAX_CONFIGURED_WORKERS,
    effective_worker_count,
    logical_cpu_count,
    probe_parallel_processes,
    worker_label,
)


RECIPE_SCHEMA = "workbench_ui_recipe/1.0.0"
DEFAULT_CONFIG = (
    "configs/V13.22_1980_recalibration_candidate_stochastic_m0_m1_20261015.json"
)
_SAFE_RECIPE_ID = re.compile(r"[^a-z0-9]+")


class WorkbenchUIError(ValueError):
    """A plain-language error suitable for display in the local UI."""


@dataclass(frozen=True)
class AnalysisDraft:
    """Serializable selections from the model-run form."""

    name: str
    config_path: str
    model: str
    p_values: tuple[float, ...] = ()
    start_year: int | None = None
    end_year: int | None = None
    trajectories: int = 1
    seeds: tuple[int, ...] = (42,)
    workers: int = 1
    parameter_overrides: Mapping[str, float] = field(default_factory=dict)
    notes: str = ""
    base_recipe_id: str | None = None

    def to_dict(self) -> dict[str, Any]:
        payload = asdict(self)
        payload["p_values"] = list(self.p_values)
        payload["seeds"] = list(self.seeds)
        payload["parameter_overrides"] = dict(self.parameter_overrides)
        return payload

    @classmethod
    def from_dict(cls, payload: Mapping[str, Any]) -> "AnalysisDraft":
        return cls(
            name=str(payload.get("name", "")),
            config_path=str(payload.get("config_path", "")),
            model=str(payload.get("model", "")),
            p_values=tuple(float(value) for value in payload.get("p_values", [])),
            start_year=_optional_int(payload.get("start_year")),
            end_year=_optional_int(payload.get("end_year")),
            trajectories=int(payload.get("trajectories", 1)),
            seeds=tuple(int(value) for value in payload.get("seeds", [])),
            workers=int(payload.get("workers", 1)),
            parameter_overrides={
                str(path): float(value)
                for path, value in dict(payload.get("parameter_overrides", {})).items()
            },
            notes=str(payload.get("notes", "")),
            base_recipe_id=(
                None
                if payload.get("base_recipe_id") in (None, "")
                else str(payload["base_recipe_id"])
            ),
        )


@dataclass(frozen=True)
class ContractControl:
    """One UI control, possibly bound to several must-equal JSON paths."""

    paths: tuple[str, ...]
    value: float
    minimum: float
    maximum: float
    unit: str
    notes: str


def _optional_int(value: Any) -> int | None:
    return None if value is None else int(value)


def parse_number_series(raw: str, *, label: str) -> tuple[float, ...]:
    """Parse comma-separated numbers or ``start:stop:step`` ranges.

    Decimal arithmetic keeps user-entered grids such as ``0:0.5:0.1`` stable.
    The range includes its stop when it lies on the declared step.
    """
    value = raw.strip()
    if not value:
        return ()
    try:
        if ":" in value:
            parts = [Decimal(item.strip()) for item in value.split(":")]
            if len(parts) != 3:
                raise WorkbenchUIError(
                    f"{label} range must use start:stop:step"
                )
            start, stop, step = parts
            if step <= 0 or stop < start:
                raise WorkbenchUIError(
                    f"{label} range needs a positive step and stop >= start"
                )
            result: list[float] = []
            current = start
            while current <= stop:
                result.append(float(current))
                if len(result) > 10_000:
                    raise WorkbenchUIError(f"{label} range is too large")
                current += step
            return tuple(result)
        return tuple(float(Decimal(item.strip())) for item in value.split(","))
    except (InvalidOperation, ValueError) as exc:
        raise WorkbenchUIError(
            f"{label} must be comma-separated numbers or start:stop:step"
        ) from exc


def parse_seed_series(raw: str) -> tuple[int, ...]:
    """Parse comma-separated uint32 seeds and inclusive integer ranges."""
    seeds: list[int] = []
    for token in raw.split(","):
        token = token.strip()
        if not token:
            continue
        if "-" in token:
            try:
                start, stop = (int(item.strip()) for item in token.split("-", 1))
            except ValueError as exc:
                raise WorkbenchUIError(f"invalid seed range: {token}") from exc
            if stop < start:
                raise WorkbenchUIError(f"descending seed range: {token}")
            seeds.extend(range(start, stop + 1))
        else:
            try:
                seeds.append(int(token))
            except ValueError as exc:
                raise WorkbenchUIError(f"invalid seed: {token}") from exc
    if not seeds:
        raise WorkbenchUIError("at least one seed is required")
    if len(seeds) != len(set(seeds)):
        raise WorkbenchUIError("seeds must be unique")
    if any(seed < 0 or seed > 2**32 - 1 for seed in seeds):
        raise WorkbenchUIError("seeds must be integers in [0, 4294967295]")
    return tuple(seeds)


def validate_draft(draft: AnalysisDraft, root: Path) -> list[str]:
    """Return all local form/configuration errors without invoking a runner."""
    errors: list[str] = []
    if draft.model not in {"M0", "M1", "both"}:
        errors.append("model must be M0, M1, or both")
    if not draft.name.strip():
        errors.append("analysis name is required")
    config = Path(draft.config_path)
    config = config if config.is_absolute() else root / config
    if not config.is_file():
        errors.append(f"configuration does not exist: {draft.config_path}")
    if draft.model in {"M1", "both"} and not draft.p_values:
        errors.append("M1 requires at least one p value")
    if draft.model == "M0" and draft.p_values:
        errors.append("p is not used by an M0-only analysis")
    if any(not math.isfinite(value) or value < 0 or value > 1 for value in draft.p_values):
        errors.append("every p value must be finite and in [0, 1]")
    if len(draft.p_values) != len(set(draft.p_values)):
        errors.append("p values must be unique")
    if draft.start_year is not None and draft.end_year is not None:
        if draft.end_year <= draft.start_year:
            errors.append("end year must be later than start year")
    if draft.trajectories < 1:
        errors.append("number of trajectories must be positive")
    if draft.trajectories > 20:
        errors.append("the current local runner supports at most 20 replicates")
    if len(draft.seeds) != draft.trajectories:
        errors.append("the number of seeds must equal the number of trajectories")
    if len(draft.seeds) != len(set(draft.seeds)):
        errors.append("seeds must be unique")
    if any(seed < 0 or seed > 2**32 - 1 for seed in draft.seeds):
        errors.append("seeds must be integers in [0, 4294967295]")
    if draft.seeds and draft.seeds != tuple(
        range(draft.seeds[0], draft.seeds[0] + len(draft.seeds))
    ):
        errors.append("the current runner requires consecutive seeds")
    if draft.workers != ALL_AVAILABLE_WORKERS and not (
        1 <= draft.workers <= MAX_CONFIGURED_WORKERS
    ):
        errors.append(
            f"workers must be 1 to {MAX_CONFIGURED_WORKERS}, or all available"
        )
    return errors


class UIWorkbenchService:
    """Repository and optional-core adapter for the browser UI."""

    def __init__(self, root: str | Path):
        self.root = Path(root).resolve()
        self.recipe_dir = self.root / "operator_configs" / "workbench_recipes"
        self._core = self._load_core()

    def _load_core(self) -> Any | None:
        try:
            from model_workbench.service import WorkbenchService
        except (ImportError, ModuleNotFoundError):
            return None
        return WorkbenchService(self.root)

    @property
    def core_available(self) -> bool:
        return self._core is not None

    def list_configurations(self) -> list[dict[str, str]]:
        registry = self.root / "CONFIG_REGISTRY.csv"
        if not registry.is_file():
            return []
        with registry.open(newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        return [
            row for row in rows
            if row.get("path", "").startswith("configs/")
            and row.get("path", "").endswith(".json")
        ]

    def configuration(self, relative_path: str) -> dict[str, Any]:
        path = Path(relative_path)
        resolved = path if path.is_absolute() else self.root / path
        try:
            return json.loads(resolved.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise WorkbenchUIError(f"cannot read configuration: {relative_path}") from exc

    def simulation_years(self, relative_path: str) -> tuple[int | None, int | None]:
        payload = self.configuration(relative_path)
        controls = payload.get("run_control", {})
        return (
            _wrapped_int(controls.get("start_year")),
            _wrapped_int(controls.get("end_year")),
        )

    def contract_controls(self, relative_path: str) -> list[ContractControl]:
        contract_path = self.root / "configs" / "collaborator_parameter_contract.json"
        if not contract_path.is_file():
            return []
        contract = json.loads(contract_path.read_text(encoding="utf-8"))
        if relative_path != contract.get("canonical_config"):
            return []
        config = self.configuration(relative_path)
        rules = {str(item["path"]): item for item in contract.get("parameters", [])}
        neighbours: dict[str, set[str]] = {path: set() for path in rules}
        for path, rule in rules.items():
            linked = str(rule.get("must_equal", ""))
            if linked:
                if linked not in rules:
                    raise WorkbenchUIError(
                        f"must_equal path is not declared in the contract: {linked}"
                    )
                neighbours[path].add(linked)
                neighbours[linked].add(path)

        seen: set[str] = set()
        controls: list[ContractControl] = []
        for path, rule in rules.items():
            if path in seen:
                continue
            pending = [path]
            component: list[str] = []
            while pending:
                member = pending.pop()
                if member in component:
                    continue
                component.append(member)
                pending.extend(sorted(neighbours[member] - set(component)))
            paths = tuple(component)
            values = [_dot_get(config, item) for item in paths]
            if len(set(float(value) for value in values)) != 1:
                raise WorkbenchUIError(f"coupled canonical values differ: {paths}")
            seen.update(paths)
            component_rules = [rules[item] for item in paths]
            units = {str(item["unit"]) for item in component_rules}
            if len(units) != 1:
                raise WorkbenchUIError(f"coupled parameter units differ: {paths}")
            minimum = max(float(item["minimum"]) for item in component_rules)
            maximum = min(float(item["maximum"]) for item in component_rules)
            if maximum < minimum:
                raise WorkbenchUIError(f"coupled parameter bounds do not overlap: {paths}")
            notes = "; ".join(
                dict.fromkeys(str(item["notes"]) for item in component_rules)
            )
            controls.append(ContractControl(
                paths=paths,
                value=float(values[0]),
                minimum=minimum,
                maximum=maximum,
                unit=units.pop(),
                notes=notes,
            ))
        return controls

    def list_recipes(self) -> list[dict[str, Any]]:
        recipes: list[dict[str, Any]] = []
        if self._core is not None:
            try:
                for item in self._core.list_recipes():
                    recipes.append(_normalise_core_recipe(item))
            except Exception as exc:  # core error should be visible, not fatal to saved recipes
                recipes.append({
                    "id": "core-error",
                    "name": "Core recipe error",
                    "source": "core",
                    "error": f"{type(exc).__name__}: {exc}",
                })
        if self.recipe_dir.is_dir():
            for path in sorted(self.recipe_dir.glob("*.json")):
                try:
                    payload = json.loads(path.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError):
                    continue
                if payload.get("schema") == RECIPE_SCHEMA:
                    recipes.append({
                        "id": payload["id"],
                        "name": payload["name"],
                        "source": "saved",
                        "path": str(path),
                    })
        return recipes

    def available_results(self) -> list[dict[str, str]]:
        """Return recognizable result sources with human-facing labels.

        Built-in examples come first, followed by completed local workbench runs
        in reverse chronological order.  The UI can still accept an explicit
        path for outputs created elsewhere.
        """
        results: list[dict[str, str]] = []
        examples = (
            (
                "examples/workbench/historical_2008",
                "Default reference fit — 2008 treatment-era change (20 runs)",
                "Lowest confirmed maximum implausibility in the retained 2008 shortlist.",
            ),
            (
                "examples/workbench/historical_2010",
                "Historical fit: 2010 treatment-era change",
                "Recorded historical calibration result; inspect only.",
            ),
        )
        for relative, label, description in examples:
            if (self.root / relative).is_dir():
                results.append({
                    "path": relative,
                    "label": label,
                    "kind": "historical",
                    "status": "retained evidence",
                    "description": description,
                })

        output_root = self.root / "outputs" / "workbench"
        manifests = sorted(
            output_root.glob("**/workbench_manifest.json"),
            key=lambda item: item.stat().st_mtime,
            reverse=True,
        ) if output_root.is_dir() else []
        for manifest_path in manifests:
            try:
                payload = json.loads(manifest_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue
            if payload.get("status") != "COMPLETE":
                continue
            plan = payload.get("resolved_plan", {})
            controls = plan.get("controls", {}) if isinstance(plan, Mapping) else {}
            model = str(controls.get("model", "model run"))
            p_value = controls.get("p")
            model_label = model if p_value is None else f"{model}, p={float(p_value):g}"
            action = str(payload.get("action", "run"))
            elapsed = payload.get("elapsed_seconds")
            elapsed_label = (
                f", {float(elapsed):.1f}s" if isinstance(elapsed, (int, float)) else ""
            )
            relative = manifest_path.parent.relative_to(self.root).as_posix()
            results.append({
                "path": relative,
                "label": f"Local {action}: {model_label}{elapsed_label}",
                "kind": "local",
                "status": "complete",
                "description": "Completed local output.",
            })
        return results

    def save_recipe(self, draft: AnalysisDraft) -> Path:
        errors = validate_draft(draft, self.root)
        if errors:
            raise WorkbenchUIError("; ".join(errors))
        recipe_id = _slug(draft.name)
        if not recipe_id:
            raise WorkbenchUIError("analysis name must contain a letter or number")
        self.recipe_dir.mkdir(parents=True, exist_ok=True)
        path = self.recipe_dir / f"{recipe_id}.json"
        if path.exists():
            raise WorkbenchUIError(
                f"a saved recipe named {recipe_id!r} already exists; choose a new analysis name"
            )
        payload = {
            "schema": RECIPE_SCHEMA,
            "id": recipe_id,
            "name": draft.name.strip(),
            "saved_utc": datetime.now(timezone.utc).isoformat(),
            "analysis": draft.to_dict(),
        }
        temporary = path.with_suffix(".json.tmp")
        temporary.write_text(
            json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8"
        )
        temporary.replace(path)
        return path

    def load_recipe(self, recipe_id: str) -> AnalysisDraft:
        path = self.recipe_dir / f"{_slug(recipe_id)}.json"
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise WorkbenchUIError(f"cannot load saved recipe: {recipe_id}") from exc
        if payload.get("schema") != RECIPE_SCHEMA:
            raise WorkbenchUIError(f"unsupported saved recipe schema: {recipe_id}")
        return AnalysisDraft.from_dict(payload["analysis"])

    def check_core_recipe(
        self, recipe_id: str, overrides: Mapping[str, Any]
    ) -> Any:
        if self._core is None:
            raise WorkbenchUIError("workbench backend adapter is not installed")
        return self._core.check(recipe_id, dict(overrides))

    def core_requests(self, draft: AnalysisDraft) -> list[dict[str, Any]]:
        """Translate a reviewed form into the allow-listed core recipe controls."""
        errors = validate_draft(draft, self.root)
        if errors:
            raise WorkbenchUIError("; ".join(errors))
        if self._core is None:
            raise WorkbenchUIError("workbench backend adapter is not installed")
        provided = dict(draft.parameter_overrides)
        controls = self.contract_controls(draft.config_path)
        allowed_paths = {path for control in controls for path in control.paths}
        unknown = set(provided) - allowed_paths
        if unknown:
            raise WorkbenchUIError(
                "saved setup contains unsupported parameter paths: "
                + ", ".join(sorted(unknown))
            )
        parameter_controls: dict[str, float] = {}
        for control in controls:
            explicit = [provided[path] for path in control.paths if path in provided]
            if explicit and any(value != explicit[0] for value in explicit[1:]):
                raise WorkbenchUIError(
                    "coupled parameter values must match: " + ", ".join(control.paths)
                )
            value = float(explicit[0] if explicit else control.value)
            if not math.isfinite(value) or not control.minimum <= value <= control.maximum:
                raise WorkbenchUIError(
                    f"{control.paths[0]} must be between {control.minimum:g} "
                    f"and {control.maximum:g}"
                )
            for path in control.paths:
                parameter_controls[f"parameter:{path}"] = value
        common = {
            "seed": draft.seeds[0],
            "replicates": draft.trajectories,
            **parameter_controls,
        }
        requests: list[dict[str, Any]] = []
        if draft.model in {"M0", "both"}:
            requests.append({
                "label": "M0",
                "recipe_id": "basic-stochastic-m0",
                "overrides": {"model": "M0", **common},
            })
        if draft.model in {"M1", "both"}:
            for value in draft.p_values:
                requests.append({
                    "label": f"M1 p={value:g}",
                    "recipe_id": "basic-stochastic-m1",
                    "overrides": {"model": "M1", "p": value, **common},
                })
        return requests

    def check_draft(self, draft: AnalysisDraft) -> list[dict[str, Any]]:
        results: list[dict[str, Any]] = []
        for request in self.core_requests(draft):
            for check in self._core.check(  # type: ignore[union-attr]
                request["recipe_id"], request["overrides"]
            ):
                payload = check.to_dict() if hasattr(check, "to_dict") else dict(check)
                results.append({"analysis": request["label"], **payload})
        return results

    def execute_draft(
        self, draft: AnalysisDraft, action: str, run_root: str | Path
    ) -> list[dict[str, Any]]:
        """Execute explicit benchmark/run plans; the caller must obtain confirmation."""
        if action not in {"benchmark", "run"}:
            raise WorkbenchUIError("action must be benchmark or run")
        base = Path(run_root).resolve()
        results: list[dict[str, Any]] = []
        for index, request in enumerate(self.core_requests(draft), start=1):
            label = _slug(request["label"]) or f"route-{index}"
            route_root = base / f"{index:02d}-{label}"
            plan = self._core.resolve(  # type: ignore[union-attr]
                request["recipe_id"], action, request["overrides"], route_root
            )
            manifest = self._core.execute(  # type: ignore[union-attr]
                plan, workers=draft.workers
            )
            result = {
                "analysis": request["label"],
                "run_root": str(route_root),
                "status": manifest.get("status", "UNKNOWN"),
                "manifest": manifest,
            }
            if action == "benchmark":
                full_plan = self._core.resolve(  # type: ignore[union-attr]
                    request["recipe_id"],
                    "run",
                    request["overrides"],
                    route_root / "estimate-only",
                )
                estimate = self._core.estimate(  # type: ignore[union-attr]
                    full_plan, manifest, workers=draft.workers
                )
                result["estimate"] = estimate.to_dict()
            results.append(result)
        return results

    def parallel_worker_check(self, workers: int) -> dict[str, Any]:
        """Test child-process launch independently of the HIV model."""

        return probe_parallel_processes(workers)

    def external_inputs_status(self) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        manifest_path = self.root / "data" / "external_inputs.json"
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise WorkbenchUIError("cannot read data/external_inputs.json") from exc
        rows: list[dict[str, Any]] = []
        for item in manifest.get("inputs", []):
            raw_path = str(item.get("path", ""))
            resolved = (self.root / raw_path).resolve()
            actual = _sha256_file(resolved) if resolved.is_file() else None
            expected = item.get("sha256")
            rows.append({
                "logical_id": item.get("logical_id", ""),
                "path": raw_path,
                "available": resolved.is_file(),
                "hash_status": (
                    "match" if actual and expected == actual
                    else "mismatch" if actual else "missing"
                ),
                "expected_sha256": expected,
                "actual_sha256": actual,
                "redistribution_status": item.get("redistribution_status", ""),
            })
        summary = {
            "path": str(manifest_path),
            "schema": manifest.get("schema"),
            "status": manifest.get("status"),
            "sha256": _sha256_file(manifest_path),
            "handoff_blockers": manifest.get("handoff_blockers", []),
        }
        return summary, rows

    def target_registers(self) -> list[str]:
        return sorted(
            path.relative_to(self.root).as_posix()
            for path in (self.root / "configs").glob("*target*json")
            if path.name != "collaborator_parameter_contract.json"
        )

    def workbook_metadata(self, path: str | Path) -> dict[str, Any]:
        workbook = Path(path)
        workbook = workbook if workbook.is_absolute() else (self.root / workbook)
        workbook = workbook.resolve()
        if not workbook.is_file():
            raise WorkbenchUIError(f"workbook is not available: {workbook}")
        try:
            with zipfile.ZipFile(workbook) as archive:
                root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
                namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                sheets = [
                    {
                        "name": node.attrib.get("name", ""),
                        "state": node.attrib.get("state", "visible"),
                    }
                    for node in root.findall("x:sheets/x:sheet", namespace)
                ]
        except (OSError, zipfile.BadZipFile, KeyError, ElementTree.ParseError) as exc:
            raise WorkbenchUIError(f"cannot inspect workbook: {workbook}") from exc
        return {
            "path": str(workbook),
            "size_bytes": workbook.stat().st_size,
            "sha256": _sha256_file(workbook),
            "sheets": sheets,
        }

    def point_patch(
        self,
        *,
        workbook_path: str | Path,
        sheet: str,
        values: Mapping[str, Any],
        note: str,
    ) -> str:
        """Return a reviewable patch document without changing source data."""
        metadata = self.workbook_metadata(workbook_path)
        sheet_names = {item["name"] for item in metadata["sheets"]}
        if sheet not in sheet_names:
            raise WorkbenchUIError(f"workbook has no sheet named {sheet!r}")
        cleaned = {
            str(key).strip(): value for key, value in values.items()
            if str(key).strip() and value not in (None, "")
        }
        if not cleaned:
            raise WorkbenchUIError("the proposed point has no values")
        payload = {
            "schema": "workbench_target_point_patch/1.0.0",
            "created_utc": datetime.now(timezone.utc).isoformat(),
            "source_workbook": metadata["path"],
            "source_workbook_sha256": metadata["sha256"],
            "operation": "append_row_for_review",
            "sheet": sheet,
            "values": cleaned,
            "note": note.strip(),
            "applied": False,
        }
        return json.dumps(payload, indent=2, ensure_ascii=False) + "\n"

    def workbook_sheet_preview(
        self, path: str | Path, sheet: str, *, max_rows: int = 200,
        max_columns: int = 60,
    ) -> dict[str, Any]:
        """Read a bounded raw worksheet preview without changing the workbook."""
        if max_rows < 1 or max_rows > 1_000 or max_columns < 1 or max_columns > 200:
            raise WorkbenchUIError("worksheet preview limits are out of range")
        metadata = self.workbook_metadata(path)
        if sheet not in {item["name"] for item in metadata["sheets"]}:
            raise WorkbenchUIError(f"workbook has no sheet named {sheet!r}")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise WorkbenchUIError(
                "openpyxl is required to preview worksheet values"
            ) from exc
        workbook = Path(path)
        workbook = workbook if workbook.is_absolute() else (self.root / workbook)
        try:
            handle = load_workbook(
                workbook.resolve(), read_only=True, data_only=True,
            )
            worksheet = handle[sheet]
            rows = [
                list(row)
                for row in worksheet.iter_rows(
                    min_row=1, max_row=max_rows, max_col=max_columns,
                    values_only=True,
                )
            ]
            worksheet_rows = worksheet.max_row
            worksheet_columns = worksheet.max_column
            handle.close()
        except (OSError, ValueError, KeyError) as exc:
            raise WorkbenchUIError(
                f"cannot preview worksheet {sheet!r}: {exc}"
            ) from exc
        while rows and all(value is None for value in rows[-1]):
            rows.pop()
        used_columns = max(
            (max((index for index, value in enumerate(row) if value is not None), default=-1) + 1
             for row in rows),
            default=0,
        )
        rows = [row[:used_columns] for row in rows]
        columns = [f"Column {index + 1}" for index in range(used_columns)]
        return {
            "sheet": sheet,
            "columns": columns,
            "rows": rows,
            "truncated": worksheet_rows > max_rows or worksheet_columns > max_columns,
            "worksheet_rows": worksheet_rows,
            "worksheet_columns": worksheet_columns,
        }

    def setup_checks(self, python_executable: str | None = None) -> list[dict[str, str]]:
        checks = [
            _file_check(self.root / "CONFIG_REGISTRY.csv", "Configuration registry"),
            _file_check(self.root / "requirements-operator.txt", "Operator requirements"),
            _file_check(self.root / DEFAULT_CONFIG, "Default M0/M1 configuration"),
        ]
        executable = python_executable or sys.executable
        checks.extend(_python_checks(executable, self.root / "requirements-operator.txt"))
        checks.append({
            "check": "Workbench backend",
            "status": "available" if self.core_available else "not available",
            "detail": (
                "model_workbench.service.WorkbenchService loaded"
                if self.core_available
                else "The interface can save drafts; validation and execution adapters are unavailable."
            ),
        })
        return checks


def _normalise_core_recipe(item: Any) -> dict[str, Any]:
    if isinstance(item, Mapping):
        recipe_id = item.get("id", item.get("recipe_id", ""))
        name = item.get("title", item.get("name", recipe_id))
    else:
        recipe_id = getattr(item, "id", getattr(item, "recipe_id", ""))
        name = getattr(item, "title", getattr(item, "name", recipe_id))
    return {"id": str(recipe_id), "name": str(name), "source": "core"}


def _slug(value: str) -> str:
    return _SAFE_RECIPE_ID.sub("-", value.strip().lower()).strip("-")[:80]


def _dot_get(payload: Mapping[str, Any], path: str) -> Any:
    current: Any = payload
    for part in path.split("."):
        if not isinstance(current, Mapping) or part not in current:
            raise WorkbenchUIError(f"configuration is missing {path}")
        current = current[part]
    return current


def _wrapped_int(value: Any) -> int | None:
    if isinstance(value, Mapping):
        value = value.get("value")
    return None if value is None else int(value)


def _file_check(path: Path, label: str) -> dict[str, str]:
    return {
        "check": label,
        "status": "available" if path.is_file() else "missing",
        "detail": str(path),
    }


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _exact_requirements(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    if not path.is_file():
        return result
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "==" not in line:
            continue
        name, version = line.split("==", 1)
        result[name.strip()] = version.strip()
    return result


def _python_checks(executable: str, requirements_path: Path) -> list[dict[str, str]]:
    script = (
        "import importlib.metadata,json,platform,sys;"
        "names=sys.argv[1:];"
        "print(json.dumps({'python':platform.python_version(),"
        "'implementation':platform.python_implementation(),"
        "'packages':{n:(importlib.metadata.version(n) if "
        "__import__('importlib.util').util.find_spec(n) else 'missing') for n in names}}))"
    )
    expected = _exact_requirements(requirements_path)
    try:
        completed = subprocess.run(
            [executable, "-c", script, *expected],
            check=False,
            capture_output=True,
            text=True,
            timeout=15,
        )
        payload = json.loads(completed.stdout) if completed.returncode == 0 else None
    except (OSError, subprocess.SubprocessError, json.JSONDecodeError):
        payload = None
    if payload is None:
        return [{
            "check": "Model Python",
            "status": "failed",
            "detail": f"Could not inspect {executable}",
        }]
    rows = [{
        "check": "Model Python",
        "status": (
            "available"
            if payload["implementation"] == "CPython"
            and payload["python"].split(".")[:2] == ["3", "11"]
            else "version mismatch"
        ),
        "detail": f"{payload['implementation']} {payload['python']} at {executable}",
    }]
    for name, version in expected.items():
        actual = payload["packages"].get(name, "missing")
        rows.append({
            "check": name,
            "status": "available" if actual == version else "version mismatch",
            "detail": f"required {version}; found {actual}",
        })
    return rows
