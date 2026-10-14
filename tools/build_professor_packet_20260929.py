#!/usr/bin/env python3
"""Build the 2026-09-29 professor-facing evidence packet from verified VM outputs."""

from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path

import numpy as np
ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "outputs/backups/conference_overnight_scaleway_20260929/restore_final_verified/outputs/suites/conference_overnight_h4d_20260929_prepared"
OUT = ROOT / "outputs/professor_evidence_package_20260929"
FIG = OUT / "figures"
FIG_V2 = OUT / "figures_v2_professor_style"

PRIMARY = "2008-break primary"
SENS = "2010-break sensitivity"
BLUE = "1F4E78"
LIGHT = "D9EAF7"
GRAY = "666666"


def rows(path: Path):
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def fnum(v):
    return float(v) if v not in (None, "") else math.nan


def selected(data, ids):
    index = {r["scenario"]: r for r in data}
    return [index[i] for i in ids if i in index]


def label(s):
    labels = {
        "base_case_context": "Oral PrEP vs no PrEP",
        "scenario_oral_matched_add_500": "Oral +500",
        "scenario_oral_matched_add_1500": "Oral +1,500",
        "scenario_oral_matched_add_3000": "Oral +3,000",
        "scenario_oral_wang_28000": "Oral to 28,000",
        "scenario_lena_substitution__switch_10": "LENA substitution 10%",
        "scenario_lena_substitution__switch_25": "LENA substitution 25%",
        "scenario_lena_substitution__switch_50": "LENA substitution 50%",
        "scenario_lena_substitution__switch_75": "LENA substitution 75%",
        "scenario_lena_substitution__switch_100": "LENA substitution 100%",
        "scenario_lena_add_on__add_on_500": "LENA add-on 500",
        "scenario_lena_add_on__add_on_1500": "LENA add-on 1,500",
        "scenario_lena_add_on__add_on_3000": "LENA add-on 3,000",
    }
    return labels.get(s, s.replace("scenario_", "").replace("_", " "))


def style_axes(ax):
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#D9D9D9", linewidth=0.7, alpha=0.7)
    ax.set_axisbelow(True)


def savefig(name):
    import matplotlib.pyplot as plt
    path = FIG / name
    plt.tight_layout()
    plt.savefig(path, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def build_figures():
    import matplotlib.pyplot as plt
    FIG.mkdir(parents=True, exist_ok=True)
    cal = rows(SRC / "calibration_comparison/case_summary.csv")
    p08 = rows(SRC / "policy_primary_2008/policy_summary.csv")
    p10 = rows(SRC / "policy_bridge_2010/policy_summary.csv")
    s08 = rows(SRC / "psa_primary_2008/psa_matrix_summary.csv")
    s10 = rows(SRC / "psa_bridge_2010/psa_matrix_summary.csv")

    # Target-by-target comparison. The archive contains an earlier legacy/delayed
    # pair and the two newly selected break-year calibrations. These are shown as
    # four available configurations, not mislabelled as a completed 2x2 factorial.
    old = rows(ROOT / "outputs/calibration/stage_diagnosis_candidate_baseline_review_20260922/target_implausibility_table.csv")
    new = rows(SRC / "calibration_comparison/target_comparison.csv")
    old_idx={(r["tuple_id"],r["target"]):r for r in old}
    targets=[r["target_key"] for r in new]
    target_labels=[t.replace("shm_new_diagnoses_msm@","Diagnoses ").replace("rivm_aids_msm@","AIDS ") for t in targets]
    observed=np.array([fnum(r["break_2008__observed_z"]) for r in new])
    series={
        "Earlier legacy": np.array([fnum(old_idx[("D_1788_v000",t)]["model_mean"]) for t in targets]),
        "Earlier delayed": np.array([fnum(old_idx[("D_1788_v109",t)]["model_mean"]) for t in targets]),
        "2008-break selected": np.array([fnum(r["break_2008__confirmation_mean"]) for r in new]),
        "2010-break selected": np.array([fnum(r["break_2010__confirmation_mean"]) for r in new]),
    }
    x=np.arange(len(targets)); offsets=[-.27,-.09,.09,.27]; markers=["o","s","^","D"]; colors=["#777777","#B07D30","#1F4E78","#2B8C7E"]
    fig,ax=plt.subplots(figsize=(11.2,5.4))
    ax.axhline(1,color="#111111",linewidth=1.1,label="Observed target")
    for (name,vals),off,m,c in zip(series.items(),offsets,markers,colors):
        ax.scatter(x+off,vals/observed,label=name,marker=m,s=40,color=c,zorder=3)
    ax.axvline(6.5,color="#C8C8C8",linewidth=1)
    ax.set_xticks(x,target_labels,rotation=50,ha="right"); ax.set_ylabel("Model mean divided by observed target")
    ax.set_title("Historical target fit across four available calibrated configurations")
    ax.grid(axis="y",color="#D9D9D9",linewidth=.7); ax.spines[["top","right"]].set_visible(False)
    ax.legend(frameon=False,ncol=3,loc="upper center",bbox_to_anchor=(.5,-.30))
    savefig("figure_1_target_fit_comparison.png")

    # 1 Calibration performance
    names = [PRIMARY, SENS]
    by_case = {r["case_id"]: r for r in cal}
    imax = [fnum(by_case["break_2008"]["confirmation_primary_i_max"]), fnum(by_case["break_2010"]["confirmation_primary_i_max"])]
    fig, ax = plt.subplots(figsize=(7.2, 3.8))
    bars = ax.barh(names, imax, color=["#1F4E78", "#7F8C8D"])
    ax.axvline(3, color="#A61C1C", linestyle="--", linewidth=1.5, label="Compatibility cutoff")
    ax.set_xlim(0, 3.3); ax.set_xlabel("Maximum confirmation implausibility")
    ax.set_title("Both calibrated structures met the prespecified cutoff")
    for b, v in zip(bars, imax): ax.text(v + .05, b.get_y()+b.get_height()/2, f"{v:.2f}", va="center")
    ax.legend(frameon=False, loc="lower right"); style_axes(ax)
    savefig("figure_1_calibration_performance.png")

    # 2 Policy infections averted
    ids = ["scenario_oral_matched_add_1500", "scenario_oral_wang_28000", "scenario_lena_substitution__switch_50", "scenario_lena_add_on__add_on_1500"]
    a, b = selected(p08, ids), selected(p10, ids)
    y = np.arange(len(ids)); h=.34
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    ax.barh(y-h/2, [fnum(r["mean_infections_averted"]) for r in a], height=h, color="#1F4E78", label=PRIMARY)
    ax.barh(y+h/2, [fnum(r["mean_infections_averted"]) for r in b], height=h, color="#9AA5AD", label=SENS)
    ax.axvline(0, color="#333333", linewidth=.8); ax.set_yticks(y, [label(i) for i in ids]); ax.invert_yaxis()
    ax.set_xlabel("Mean infections averted, 15 paired seeds"); ax.set_title("Projected effects depend materially on epidemic structure")
    ax.legend(frameon=False, loc="lower right"); style_axes(ax)
    savefig("figure_2_policy_infections_averted.png")

    # 3 Policy NMB
    ids = ["scenario_oral_matched_add_500", "scenario_oral_matched_add_1500", "scenario_oral_matched_add_3000", "scenario_lena_substitution__switch_10", "scenario_lena_substitution__switch_50", "scenario_lena_substitution__switch_100", "scenario_lena_add_on__add_on_500", "scenario_lena_add_on__add_on_1500", "scenario_lena_add_on__add_on_3000"]
    a = selected(p08, ids); b = selected(p10, ids); y=np.arange(len(ids)); h=.34
    fig, ax=plt.subplots(figsize=(9.0, 6.0))
    ax.barh(y-h/2, [fnum(r["mean_nmb"])/1e9 for r in a], height=h, color="#1F4E78", label=PRIMARY)
    # Missing scenarios in sensitivity are intentionally absent; align through lookup.
    bidx={r["scenario"]:r for r in p10}
    ax.barh(y+h/2, [fnum(bidx[i]["mean_nmb"])/1e9 if i in bidx else np.nan for i in ids], height=h, color="#9AA5AD", label=SENS)
    ax.axvline(0,color="#A61C1C",linewidth=1.2); ax.set_yticks(y,[label(i) for i in ids]); ax.invert_yaxis()
    ax.set_xlabel("Mean net monetary benefit, EUR billions (WTP EUR 80,000/QALY)")
    ax.set_title("All tested lenacapavir policies had negative mean NMB")
    ax.legend(frameon=False,loc="lower right"); style_axes(ax)
    savefig("figure_3_policy_nmb.png")

    # 4 PSA decision uncertainty
    ids=["scenario_oral_matched_add_1500","scenario_oral_wang_28000","scenario_lena_substitution__switch_50","scenario_lena_add_on__add_on_1500"]
    a=selected(s08,ids); b=selected(s10,ids); y=np.arange(len(ids)); h=.34
    fig,ax=plt.subplots(figsize=(8.4,4.6))
    ax.barh(y-h/2,[100*fnum(r["probability_cost_effective"]) for r in a],height=h,color="#1F4E78",label=PRIMARY)
    ax.barh(y+h/2,[100*fnum(r["probability_cost_effective"]) for r in b],height=h,color="#9AA5AD",label=SENS)
    ax.set_xlim(0,40); ax.set_yticks(y,[label(i) for i in ids]); ax.invert_yaxis(); ax.set_xlabel("Probability cost-effective (%)")
    ax.set_title("PSA uncertainty does not make lenacapavir the favored option")
    ax.legend(frameon=False,loc="lower right"); style_axes(ax)
    savefig("figure_4_psa_probability_cost_effective.png")

    # 5 PSA QALY uncertainty
    ids=["scenario_oral_matched_add_1500","scenario_oral_wang_28000","scenario_lena_substitution__switch_50","scenario_lena_add_on__add_on_1500"]
    a=selected(s08,ids); y=np.arange(len(ids))
    means=np.array([fnum(r["mean_delta_qaly"]) for r in a]); lo=np.array([fnum(r["p025_delta_qaly"]) for r in a]); hi=np.array([fnum(r["p975_delta_qaly"]) for r in a])
    fig,ax=plt.subplots(figsize=(8.4,4.6))
    ax.errorbar(means,y,xerr=[means-lo,hi-means],fmt="o",color="#1F4E78",ecolor="#7FA7C4",capsize=4)
    ax.axvline(0,color="#A61C1C",linewidth=1.1); ax.set_yticks(y,[label(i) for i in ids]); ax.invert_yaxis(); ax.set_xlabel("Incremental QALYs, mean and 95% percentile interval")
    ax.set_title("Primary-model PSA shows wide health-outcome uncertainty")
    style_axes(ax); savefig("figure_5_psa_qaly_uncertainty.png")


def build_figures_v2():
    """Compact, professor-packet-style replacement figures."""
    import matplotlib.pyplot as plt
    global FIG
    old_fig = FIG
    FIG = FIG_V2
    FIG.mkdir(parents=True, exist_ok=True)
    old = rows(ROOT / "outputs/calibration/stage_diagnosis_candidate_baseline_review_20260922/target_implausibility_table.csv")
    new = rows(SRC / "calibration_comparison/target_comparison.csv")
    old_idx={(r["tuple_id"],r["target"]):r for r in old}; targets=[r["target_key"] for r in new]
    configs={
        "Earlier legacy": [fnum(old_idx[("D_1788_v000",t)]["model_mean"]) for t in targets],
        "Earlier delayed diagnosis": [fnum(old_idx[("D_1788_v109",t)]["model_mean"]) for t in targets],
        "2008 break selected": [fnum(r["break_2008__confirmation_mean"]) for r in new],
        "2010 break selected": [fnum(r["break_2010__confirmation_mean"]) for r in new],
    }
    observed=np.array([fnum(r["break_2008__observed_z"]) for r in new])
    diag=np.arange(7); aids=np.arange(7,13); years_diag=[int(t.split("@")[-1]) for t in targets[:7]]; years_aids=[int(t.split("@")[-1]) for t in targets[7:]]
    fig,axs=plt.subplots(2,2,figsize=(10.5,6.5),sharey=False)
    colors=["#315B7D","#D8892B","#315B7D","#D8892B"]
    markers=["o","s","^","D"]
    for ax,(name,vals),c,m in zip(axs.ravel(),configs.items(),colors,markers):
        vals=np.array(vals)
        ax.plot(years_diag,observed[diag],"-o",color="#111111",lw=1.6,ms=4,label="Observed diagnoses")
        ax.plot(years_diag,vals[diag],f"-{m}",color=c,lw=1.8,ms=4,label="Model diagnoses")
        ax.plot(years_aids,observed[aids],"--o",color="#555555",lw=1.3,ms=4,label="Observed AIDS")
        ax.plot(years_aids,vals[aids],f"--{m}",color=c,lw=1.5,ms=4,label="Model AIDS")
        ax.set_title(name,loc="left",fontsize=10,fontweight="bold"); ax.grid(axis="y",color="#DDDDDD",lw=.6); ax.spines[["top","right"]].set_visible(False); ax.tick_params(labelsize=8)
    axs[0,0].set_ylabel("Annual count"); axs[1,0].set_ylabel("Annual count")
    handles,labels=axs[0,0].get_legend_handles_labels(); fig.legend(handles,labels,ncol=4,frameon=False,loc="lower center",fontsize=8)
    fig.suptitle("Observed calibration targets and model means",fontsize=13,fontweight="bold",x=.06,ha="left")
    plt.tight_layout(rect=(0,.07,1,.95)); plt.savefig(FIG/"figure_1_calibration_small_multiples.png",dpi=260,bbox_inches="tight",facecolor="white"); plt.close()

    # Direct target residual chart, easier to see over/under-prediction.
    fig,ax=plt.subplots(figsize=(10.5,4.8)); x=np.arange(13); offs=[-.27,-.09,.09,.27]
    for (name,vals),off,c,m in zip(configs.items(),offs,colors,markers): ax.scatter(x+off,(np.array(vals)-observed)/observed*100,label=name,color=c,marker=m,s=34)
    ax.axhline(0,color="#111",lw=1); ax.axvline(6.5,color="#AAAAAA",lw=.8)
    labs=[t.replace("shm_new_diagnoses_msm@","Dx ").replace("rivm_aids_msm@","AIDS ") for t in targets]
    ax.set_xticks(x,labs,rotation=45,ha="right",fontsize=8); ax.set_ylabel("Difference from observed target (%)"); ax.set_title("Relative calibration error by target year",loc="left",fontweight="bold")
    ax.grid(axis="y",color="#DDDDDD",lw=.6); ax.spines[["top","right"]].set_visible(False); ax.legend(frameon=False,ncol=2,fontsize=8,loc="upper left")
    savefig("figure_2_calibration_relative_error.png")

    p08=rows(SRC/"policy_primary_2008/policy_summary.csv"); p10=rows(SRC/"policy_bridge_2010/policy_summary.csv")
    ids=["scenario_lena_substitution__switch_10","scenario_lena_substitution__switch_25","scenario_lena_substitution__switch_50","scenario_lena_substitution__switch_75","scenario_lena_substitution__switch_100","scenario_lena_add_on__add_on_500","scenario_lena_add_on__add_on_1500","scenario_lena_add_on__add_on_3000"]
    idx08={r["scenario"]:r for r in p08}; idx10={r["scenario"]:r for r in p10}; labs=[label(i).replace("LENA ","") for i in ids]
    fig,axs=plt.subplots(1,2,figsize=(10.5,4.4),sharey=True)
    for ax,title,idx in [(axs[0],"2008 break primary",idx08),(axs[1],"2010 break sensitivity",idx10)]:
        vals=[fnum(idx[i]["mean_infections_averted"]) if i in idx else np.nan for i in ids]; cols=["#315B7D"]*5+["#2E8B57"]*3
        ax.barh(np.arange(8),vals,color=cols); ax.axvline(0,color="#222",lw=.8); ax.set_title(title,fontsize=10,fontweight="bold"); ax.grid(axis="x",color="#DDD",lw=.6); ax.spines[["top","right"]].set_visible(False)
    axs[0].set_yticks(np.arange(8),labs,fontsize=8); axs[0].invert_yaxis(); axs[0].set_xlabel("Mean infections averted"); axs[1].set_xlabel("Mean infections averted")
    fig.suptitle("Infections averted across lenacapavir deployment scenarios",fontsize=13,fontweight="bold",x=.06,ha="left"); plt.tight_layout(rect=(0,0,1,.94)); savefig("figure_3_infections_averted.png")

    s08=rows(SRC/"psa_primary_2008/psa_matrix_summary.csv"); s10=rows(SRC/"psa_bridge_2010/psa_matrix_summary.csv"); ids=["scenario_oral_matched_add_1500","scenario_oral_wang_28000","scenario_lena_substitution__switch_50","scenario_lena_add_on__add_on_1500"]
    a={r["scenario"]:r for r in s08}; b={r["scenario"]:r for r in s10}; y=np.arange(4)
    fig,ax=plt.subplots(figsize=(9.2,4.2)); ax.scatter([fnum(a[i]["mean_nmb"])/1e6 for i in ids],y-.12,color="#315B7D",s=52,label="2008 break primary"); ax.scatter([fnum(b[i]["mean_nmb"])/1e6 for i in ids],y+.12,color="#D8892B",marker="s",s=48,label="2010 break sensitivity")
    ax.axvline(0,color="#B22222",ls="--",lw=1.2); ax.set_yticks(y,[label(i) for i in ids]); ax.invert_yaxis(); ax.set_xlabel("Mean net monetary benefit (EUR million)"); ax.set_title("Probabilistic economic results",loc="left",fontweight="bold")
    ax.grid(axis="x",color="#DDD",lw=.6); ax.spines[["top","right"]].set_visible(False); ax.legend(frameon=False,fontsize=8); savefig("figure_4_psa_nmb.png")
    FIG=old_fig


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def set_cell(cell, text, bold=False, color=None, size=8.5):
    cell.text=""; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(str(text)); r.bold=bold; r.font.name="Arial"; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, body, widths=None):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style="Table Grid"
    for i,h in enumerate(headers): shade(table.rows[0].cells[i], BLUE); set_cell(table.rows[0].cells[i],h,True,"FFFFFF",8)
    for ri,row in enumerate(body):
        cells=table.add_row().cells
        for i,v in enumerate(row):
            if ri%2: shade(cells[i],"F3F6F8")
            set_cell(cells[i],v,size=8)
    return table


def add_heading(doc, text, level=1):
    p=doc.add_heading(text,level=level); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5)
    for r in p.runs: r.font.name="Arial"; r.font.color.rgb=RGBColor(0,0,0)
    return p


def add_para(doc,text,bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.08
    if bold_lead and text.startswith(bold_lead):
        r=p.add_run(bold_lead); r.bold=True; text=text[len(bold_lead):]
    p.add_run(text)
    for r in p.runs: r.font.name="Arial"; r.font.size=Pt(9.5)
    return p


def add_figure(doc,path,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(6.55))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(7)
    r=c.add_run(caption); r.bold=True; r.font.name="Arial"; r.font.size=Pt(8.5)


def build_docx():
    global Document, WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT, WD_ALIGN_PARAGRAPH
    global OxmlElement, qn, Inches, Pt, RGBColor
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    p08=rows(SRC/"policy_primary_2008/policy_summary.csv"); p10=rows(SRC/"policy_bridge_2010/policy_summary.csv")
    s08=rows(SRC/"psa_primary_2008/psa_matrix_summary.csv"); s10=rows(SRC/"psa_bridge_2010/psa_matrix_summary.csv")
    cal=rows(SRC/"calibration_comparison/case_summary.csv"); c={r["case_id"]:r for r in cal}
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    styles=doc.styles
    styles["Normal"].font.name="Arial"; styles["Normal"].font.size=Pt(9.5)
    styles["Title"].font.name="Arial"; styles["Title"].font.size=Pt(24); styles["Title"].font.color.rgb=RGBColor(0,0,0)
    title=doc.add_paragraph(style="Title"); title.add_run("Updated Model Run Evidence Summary")
    sub=doc.add_paragraph("Professor review packet  |  Verified VM run completed 29 September 2026")
    sub.runs[0].font.name="Arial"; sub.runs[0].font.size=Pt(10); sub.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
    add_para(doc,"Bottom line. The overnight suite completed and passed its validation gates. Both epidemic structures met the calibration compatibility criterion. The 2010-break structure had the lower maximum implausibility, but the 2008-break structure had the better overall confirmation objective and remains the primary analysis. At the configured lenacapavir price, every tested lenacapavir policy had negative mean net monetary benefit in the 15-seed policy analysis. The 1,000-draw PSA introduced wide uncertainty but did not reverse the mean economic conclusion.","Bottom line.")

    add_heading(doc,"Run completion and calibration",1)
    body=[
        [PRIMARY,c["break_2008"]["selected_candidate_id"],"Yes",f'{fnum(c["break_2008"]["confirmation_primary_i_max"]):.2f}',f'{fnum(c["break_2008"]["confirmation_j"]):.2f}',c["break_2008"]["n_confirmation_primary_pass"]],
        [SENS,c["break_2010"]["selected_candidate_id"],"Yes",f'{fnum(c["break_2010"]["confirmation_primary_i_max"]):.2f}',f'{fnum(c["break_2010"]["confirmation_j"]):.2f}',c["break_2010"]["n_confirmation_primary_pass"]],
    ]
    add_table(doc,["Structure","Candidate","Primary pass","Maximum I","Objective J","Passing candidates"],body)
    add_para(doc,"Both selected candidates were below the frozen maximum-implausibility cutoff of 3. The 2010-break candidate had the lower maximum I, while the 2008-break candidate had the lower overall objective. This supports treating structural timing as a sensitivity rather than claiming one structure is validated as true.")
    add_figure(doc,FIG/"figure_1_target_fit_comparison.png","Figure 1  Model means relative to all 13 observed calibration targets")
    add_para(doc,"Figure 1 places all 13 count targets on a common ratio scale; 1.0 is the observed target. It includes the earlier legacy and delayed-diagnosis configurations plus the two newly selected break-year calibrations. These are the four available calibrated outputs, not a completed legacy-versus-delayed factorial within each break year. The latter would require two additional matched runs.")
    add_figure(doc,FIG/"figure_1_calibration_performance.png","Figure 2  Maximum implausibility for the newly selected break-year structures")

    add_heading(doc,"Policy analysis",1)
    add_para(doc,"The policy analysis used 15 paired common-random-number seeds at fixed parameter values. It characterizes stochastic simulation uncertainty, not parameter uncertainty. Net monetary benefit is reported at EUR 80,000 per QALY.")
    ids=["scenario_oral_matched_add_1500","scenario_oral_wang_28000","scenario_lena_substitution__switch_50","scenario_lena_add_on__add_on_1500"]
    a={r["scenario"]:r for r in p08}; b={r["scenario"]:r for r in p10}; body=[]
    for i in ids:
        for struct,x in [(PRIMARY,a.get(i)),(SENS,b.get(i))]:
            if x: body.append([struct,label(i),f'{fnum(x["mean_infections_averted"]):,.1f}',f'{fnum(x["mean_delta_qaly"]):,.1f}',f'{fnum(x["mean_delta_cost"])/1e6:,.1f}',f'{fnum(x["mean_nmb"])/1e6:,.1f}',f'{100*fnum(x["probability_positive_nmb"]):.0f}%'])
    add_table(doc,["Structure","Scenario","Infections averted","Delta QALYs","Delta cost EUR m","NMB EUR m","P NMB > 0"],body)
    add_para(doc,"For the primary model, 50% substitution averted 15.7 infections and the 1,500-user add-on averted 78.0; their mean NMB values were approximately EUR -1.27 billion and EUR -1.16 billion. The full substitution and 3,000-user add-on scenarios were still less favorable economically. Persistence sensitivities also remained mean-NMB negative.")
    add_figure(doc,FIG/"figure_2_policy_infections_averted.png","Figure 3  Mean infections averted in selected policy scenarios")
    add_figure(doc,FIG/"figure_3_policy_nmb.png","Figure 4  Mean policy-analysis net monetary benefit")

    add_heading(doc,"Probabilistic sensitivity analysis",1)
    add_para(doc,"The PSA used 1,000 matched draws for each structure. These results should not be combined numerically with the 15-seed policy intervals: the PSA varies uncertain inputs, whereas the policy analysis holds parameters fixed and varies simulation seeds.")
    ai={r["scenario"]:r for r in s08}; bi={r["scenario"]:r for r in s10}; body=[]
    for i in ids:
        for struct,x in [(PRIMARY,ai.get(i)),(SENS,bi.get(i))]:
            if x: body.append([struct,label(i),x["draws"],f'{fnum(x["mean_infections_averted"]):,.1f}',f'{fnum(x["mean_delta_qaly"]):,.1f}',f'{fnum(x["mean_nmb"])/1e6:,.1f}',f'{100*fnum(x["probability_cost_effective"]):.1f}%'])
    add_table(doc,["Structure","Scenario","Draws","Infections averted","Delta QALYs","NMB EUR m","P cost-effective"],body)
    add_para(doc,"In the primary structure, the 1,500-user lenacapavir add-on had mean NMB of approximately EUR -709 million and a 10.9% probability of cost-effectiveness. The corresponding 2010-break result was approximately EUR -221 million with a 19.4% probability. The wide QALY intervals, including sign changes in several scenarios, show that the health result is structurally and parametrically uncertain even though mean acquisition costs dominate the economic conclusion.")
    add_figure(doc,FIG/"figure_4_psa_probability_cost_effective.png","Figure 5  Probability of cost-effectiveness in the 1,000-draw PSA")
    add_figure(doc,FIG/"figure_5_psa_qaly_uncertainty.png","Figure 6  Primary-model PSA uncertainty in incremental QALYs")

    add_heading(doc,"Interpretation for the thesis",1)
    for text in [
        "The run is technically usable. The suite status is complete, both bound models were confirmed, policy and PSA outputs were finalized for both structures, paired structural outputs were created, and the validation record reports no retained failure artifacts.",
        "The substantive conclusion is narrower than a simple effectiveness claim. Lenacapavir can avert infections in many scenarios, but the modeled health effect is uncertain and the configured acquisition cost produces negative mean NMB across all tested lenacapavir strategies.",
        "Epidemic structure materially changes benefit estimates. The 2008-break model should therefore remain the primary analysis and the 2010-break model should be reported alongside it as a structural sensitivity.",
        "Oral-PrEP expansion is the more favorable comparator family. In the primary PSA, expansion by 3,000 users had positive mean NMB of EUR 11.7 million, but only a 30.1% probability of cost-effectiveness; this is suggestive rather than decisive.",
    ]: add_para(doc,text)

    add_heading(doc,"Limitations and reporting rules",1)
    for t in [
        "Calibration compatibility does not establish predictive validity or identify the true epidemic structure.",
        "The 15-seed policy analysis measures stochastic simulation uncertainty at fixed inputs. The 1,000-draw PSA is the parameter-uncertainty analysis.",
        "Mean NMB and cost-effectiveness probability are decision summaries, not proof that an intervention is or is not effective in every draw.",
        "Results are conditional on the configured scenarios, 2025 to 2064 horizon, model assumptions, and EUR 80,000 per-QALY decision threshold used in these outputs.",
        "The full verification workbook preserves summary, seed-level, draw-level, structural-comparison, and provenance records.",
    ]:
        p=doc.add_paragraph(style=None); p.style=doc.styles["List Bullet"]; p.add_run(t)
        for r in p.runs: r.font.name="Arial"; r.font.size=Pt(9)

    add_heading(doc,"Evidence files",1)
    add_para(doc,"The accompanying workbook, Updated Model Run Verification.xlsx, contains the calibration selection, policy summaries, PSA summaries, paired structural comparisons, seed-level policy records, draw-level PSA records, convergence diagnostics, and provenance paths used in this report.")
    core=doc.core_properties; core.title="Updated Model Run Evidence Summary"; core.subject="Verified 2026-09-29 model run professor packet"; core.author="Rishi Malhotra"
    OUT.mkdir(parents=True,exist_ok=True)
    doc.save(OUT/"Updated_Model_Run_Evidence_Summary.docx")


def build_json():
    datasets={
        "calibration": SRC/"calibration_comparison/case_summary.csv",
        "policy_2008": SRC/"policy_primary_2008/policy_summary.csv",
        "policy_2010": SRC/"policy_bridge_2010/policy_summary.csv",
        "psa_2008": SRC/"psa_primary_2008/psa_matrix_summary.csv",
        "psa_2010": SRC/"psa_bridge_2010/psa_matrix_summary.csv",
        "policy_seed_2008": SRC/"policy_primary_2008/policy_seed_results.csv",
        "policy_seed_2010": SRC/"policy_bridge_2010/policy_seed_results.csv",
        "psa_draw_2008": SRC/"psa_primary_2008/psa_matrix_incremental.csv",
        "psa_draw_2010": SRC/"psa_bridge_2010/psa_matrix_incremental.csv",
        "paired_policy": SRC/"combined/paired_policy_structural_summary.csv",
        "paired_psa": SRC/"combined/paired_psa_structural_summary.csv",
        "convergence_2008": SRC/"psa_primary_2008/psa_matrix_convergence.csv",
        "convergence_2010": SRC/"psa_bridge_2010/psa_matrix_convergence.csv",
        "earlier_target_fit": ROOT/"outputs/calibration/stage_diagnosis_candidate_baseline_review_20260922/target_implausibility_table.csv",
        "new_target_fit": SRC/"calibration_comparison/target_comparison.csv",
    }
    payload={k:{"source":str(v.relative_to(ROOT)),"rows":rows(v)} for k,v in datasets.items()}
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/"packet_data.json").write_text(json.dumps(payload),encoding="utf-8")


if __name__ == "__main__":
    ap=argparse.ArgumentParser(); ap.add_argument("mode",choices=["figures","figures_v2","docx","json","all"]); args=ap.parse_args()
    if args.mode in ("figures","all"): build_figures()
    if args.mode in ("figures_v2","all"): build_figures_v2()
    if args.mode in ("json","all"): build_json()
    if args.mode in ("docx","all"): build_docx()
