#!/usr/bin/env python3
"""Create a compact professor-style copy of the updated evidence packet."""

from __future__ import annotations

import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PKG = ROOT / "outputs/professor_evidence_package_20260929"
SOURCE = PKG / "Updated_Model_Run_Evidence_Summary.docx"
OUTPUT = PKG / "Updated_Model_Runs_and_Results_Professor_Style.docx"
FIG = PKG / "figures_v2_professor_style"
CHARCOAL = "404040"
ALT = "F2F2F2"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def replace_media(docx_path: Path):
    replacements = {
        "word/media/image1.png": FIG / "figure_1_calibration_small_multiples.png",
        "word/media/image2.png": FIG / "figure_2_calibration_relative_error.png",
        "word/media/image3.png": FIG / "figure_3_infections_averted.png",
        "word/media/image5.png": FIG / "figure_4_psa_nmb.png",
    }
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        with zipfile.ZipFile(docx_path) as zin:
            zin.extractall(tmp)
        for member, image in replacements.items():
            shutil.copyfile(image, tmp / member)
        rebuilt = tmp / "rebuilt.docx"
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as zout:
            for path in sorted(tmp.rglob("*")):
                if path.is_file() and path != rebuilt:
                    zout.write(path, path.relative_to(tmp).as_posix())
        shutil.copyfile(rebuilt, docx_path)


def main():
    doc = Document(SOURCE)
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.6)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.02
    for name, size in [("Title", 14), ("Heading 1", 11), ("Heading 2", 9.5)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(3)

    replacements = {
        "Updated Model Run Evidence Summary": "Updated Model Runs and Results",
        "Figure 1  Model means relative to all 13 observed calibration targets": "Figure 1  Observed calibration targets and model means",
        "Figure 1 places all 13 count targets on a common ratio scale; 1.0 is the observed target. It includes the earlier legacy and delayed-diagnosis configurations plus the two newly selected break-year structures. These are four available calibrated configurations, not a completed 2 × 2 factorial.": "Figure 1 compares observed diagnoses and AIDS targets with each of the four available calibrated configurations. The earlier legacy and delayed-diagnosis models and the new 2008- and 2010-break calibrations are separate analyses; they are not a complete 2 × 2 factorial.",
        "Figure 2  Maximum implausibility for the newly selected break-year structures": "Figure 2  Relative calibration error by target year",
        "Figure 3  Mean infections averted in selected policy scenarios": "Figure 3  Infections averted across lenacapavir deployment scenarios",
        "Figure 5  Probability of cost-effectiveness in the 1,000-draw PSA": "Figure 5  Mean net monetary benefit in the 1,000-draw PSA",
        "Updated Model Run Verification.xlsx": "Updated Model Run Verification — Professor Style.xlsx",
    }
    for p in doc.paragraphs:
        if p.text in replacements:
            p.text = replacements[p.text]
        if p.text.startswith("Figure 1 places all 13 count targets"):
            p.text = ("Figure 1 compares observed diagnoses and AIDS targets with each of the four available "
                      "calibrated configurations. The earlier legacy and delayed-diagnosis models and the new "
                      "2008- and 2010-break calibrations are separate analyses; they are not a complete 2 × 2 factorial.")
        if p.text.startswith("The accompanying workbook,"):
            p.text = ("The accompanying workbook, Updated Model Run Verification — Professor Style.xlsx, contains "
                      "the calibration selection, target-by-target comparison, policy summaries, PSA summaries, "
                      "paired structural comparisons, seed-level policy records, draw-level PSA records, convergence "
                      "diagnostics, and provenance paths used in this report.")
        if p.text.startswith("Figure "):
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                run.font.name = "Arial"; run.font.size = Pt(7.5); run.font.bold = True

    # Compact title and subtitle.
    doc.paragraphs[0].style = styles["Title"]
    doc.paragraphs[0].paragraph_format.space_after = Pt(1)
    p_pr = doc.paragraphs[0]._p.get_or_add_pPr()
    for border in p_pr.findall(qn("w:pBdr")):
        p_pr.remove(border)
    doc.paragraphs[1].paragraph_format.space_after = Pt(5)
    for run in doc.paragraphs[1].runs:
        run.font.name = "Arial"; run.font.size = Pt(8); run.font.italic = True; run.font.color.rgb = RGBColor(90, 90, 90)

    # Reformat every table in the same compact grayscale style as the original packet.
    for table in doc.tables:
        table.autofit = True
        set_repeat_table_header(table.rows[0])
        for ci, cell in enumerate(table.rows[0].cells):
            shade(cell, CHARCOAL)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Arial"; run.font.size = Pt(7.2); run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)
        for ri, row in enumerate(table.rows[1:], 1):
            if ri % 2 == 0:
                for cell in row.cells: shade(cell, ALT)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        run.font.name = "Arial"; run.font.size = Pt(7.1); run.font.color.rgb = RGBColor(25,25,25)

    # Fit the four replacement graphics to the usable page width without distortion.
    ratios = [10.5/6.5, 10.5/4.8, 10.5/4.4, 9.2/4.2]
    for shape_index, ratio in zip([0, 1, 2, 4], ratios):
        shape = doc.inline_shapes[shape_index]
        shape.width = Inches(6.9)
        shape.height = Inches(6.9 / ratio)

    doc.save(OUTPUT)
    replace_media(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
