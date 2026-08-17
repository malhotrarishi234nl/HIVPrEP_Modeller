import csv
import os

from openpyxl import load_workbook

from hta_report import HTAReportWriter
from scenario_comparator import ScenarioComparator, ScenarioResult


def _cost_breakdown(total, prep_oral=0.0, prep_lena=0.0, hiv_art=0.0):
    monitoring = total * 0.2
    hiv_care = total - prep_oral - prep_lena - hiv_art - monitoring
    return {
        "drug": prep_oral + prep_lena + hiv_art,
        "prep_drug": prep_oral + prep_lena,
        "prep_drug_oral": prep_oral,
        "prep_drug_lenacapavir": prep_lena,
        "hiv_art_drug": hiv_art,
        "monitoring": monitoring,
        "hiv_care": hiv_care,
        "adverse_events": 0.0,
        "terminal_care": 0.0,
    }


def _result(arm, scenario, cost, qalys, infections):
    if arm in ("oral_prep", "oral_prep_expansion"):
        breakdown = _cost_breakdown(cost, prep_oral=400.0, hiv_art=200.0)
    elif arm == "lenacapavir":
        breakdown = _cost_breakdown(cost, prep_lena=700.0, hiv_art=200.0)
    else:
        breakdown = _cost_breakdown(cost, hiv_art=600.0)

    return ScenarioResult(
        name=arm,
        analysis_scenario=scenario,
        total_qalys_disc=qalys,
        total_qalys_undisc=qalys,
        total_costs_disc=cost,
        total_costs_undisc=cost,
        cost_breakdown_disc=breakdown,
        cumul_new_infections=infections,
        timestep_records=[
            {
                "step": 1,
                "year": 2025.0,
                "step_fraction": 1.0,
                "total_population": 1000,
                "step_qalys_disc": qalys,
                "step_qalys_undisc": qalys,
                "step_costs_disc": cost,
                "step_costs_undisc": cost,
                "step_new_infections": infections,
                "discount_factor_health": 1.0,
                "discount_factor_cost": 1.0,
            }
        ],
        simulation_records=[
            {
                "step": idx,
                "year": float(year),
                "total_population": 1000 + idx * 10,
                "msm_total": 500,
                "msm_susceptible": 480,
                "msm_eligible_prep": 200,
                "msm_high_risk": 100,
                "msm_high_risk_susceptible": 90,
                "msm_high_risk_on_prep": 20,
                "msm_hiv_pos": 20,
                "adult_hiv_pos": 25 + idx,
                "prevalence": 0.02 + idx / 10000,
                "hiv_deaths": idx,
                "msm_on_oral": 25 if arm in ("oral_prep", "oral_prep_expansion") else 0,
                "msm_on_lena": 10 if arm == "lenacapavir" else 0,
                "prep_oral_initiations": 2 if arm in ("oral_prep", "oral_prep_expansion") else 0,
                "prep_lena_initiations": 1 if arm == "lenacapavir" else 0,
                "prep_oral_dropouts": 0,
                "prep_lena_dropouts": 0,
                "new_infections": infections + idx,
                "new_infections_msm": infections + idx,
                "new_infections_msm_on_prep": 0,
            }
            for idx, year in enumerate((2025, 2030, 2031), start=1)
        ],
    )


def _split_runtime_config():
    return {
        "outputs": {
            "excel_report": {
                "enabled": True,
                "detail_sheet_mode": "off",
                "include_full_data_tabs": False,
                "include_quarterly_timestep_tabs": False,
                "include_subgroup_detail_tabs": False,
            },
            "baseline_full_tab": {"enabled": False},
            "scenario_full_data_tabs": {"enabled": False},
            "transparency_tabs": {"enabled": True},
        },
        "comparison_selection": {
            "primary_comparison_rule": "exactly_one_required",
            "secondary_comparisons": "all_decision_relevant_non_primary",
        },
        "base_cases": {
            "oral_prep_standard_of_care": {
                "label": "Oral PrEP standard of care",
                "include_in_decision_summary": True,
            },
            "no_prep_counterfactual": {
                "label": "No PrEP counterfactual",
                "include_in_decision_summary": False,
            },
        },
        "scenarios": {
            "scenario_a_lena_substitution": {
                "label": "Lena substitution",
                "primary_comparison": True,
            },
            "scenario_b_lena_add_on": {
                "label": "Lena add-on",
                "primary_comparison": False,
            },
        },
    }


def _hta_config():
    runtime_config = _split_runtime_config()
    return {
        "_format_version": "v2",
        "_runtime_config": runtime_config,
        "outputs": runtime_config["outputs"],
        "wtp_threshold_eur_per_qaly": {"value": 80000},
        "time_horizon": {"freeze_year": 2025, "evaluation_years": 1},
        "discounting": {
            "health_effects_rate": {"value": 0.015},
            "cost_rate": {"value": 0.04},
            "base_year": {"value": 2025},
        },
        "perspective": {"primary": "payer"},
        "prevention_costs": {
            "arm_specific": {
                "lenacapavir": {
                    "injection_subsequent": {"value": 20000},
                },
            },
        },
        "adverse_events": {
            "items": [
                {
                    "id": "lenacapavir_isr_ae",
                    "label": "Transient injection site reactions",
                    "prevalence": {"value": 0.405, "source": "test ISR source"},
                    "utility_decrement": {"value": -0.01, "source": "test utility source"},
                    "cost": {"value": 0, "source": "test zero-cost source"},
                }
            ]
        },
        "analysis_scenarios": [
            {
                "name": "scenario_a_lena_substitution__switch_50",
                "label": "Lena substitution, 50% switch",
                "comparison_type": "substitution",
                "comparator_arm": "oral_prep",
                "intervention_arm": "lenacapavir",
                "primary_comparison": True,
                "price_sensitivity": {
                    "enabled": True,
                    "percent_values": [0.5, 1.0],
                },
            },
            {
                "name": "scenario_b_lena_add_on__add_on_500",
                "label": "Lena add-on, 500 users",
                "comparison_type": "add_on",
                "comparator_arm": "oral_prep",
                "intervention_arm": "lenacapavir",
                "primary_comparison": False,
            },
            {
                "name": "scenario_c_oral_expansion_benchmark",
                "label": "Oral PrEP expansion benchmark",
                "comparison_type": "oral_expansion_benchmark",
                "comparator_arm": "oral_prep",
                "intervention_arm": "oral_prep_expansion",
                "primary_comparison": False,
            },
        ],
    }


def _all_results():
    scenarios = {
        "base_case_context": {
            "no_prep": _result(
                "no_prep", "base_case_context", cost=1000.0, qalys=100.0, infections=10
            ),
            "oral_prep": _result(
                "oral_prep", "base_case_context", cost=1200.0, qalys=102.0, infections=8
            ),
        }
    }
    for scenario in [
        "scenario_a_lena_substitution__switch_50",
        "scenario_b_lena_add_on__add_on_500",
    ]:
        scenarios[scenario] = {
            "no_prep": _result("no_prep", scenario, cost=1000.0, qalys=100.0, infections=10),
            "oral_prep": _result("oral_prep", scenario, cost=1200.0, qalys=102.0, infections=8),
            "lenacapavir": _result("lenacapavir", scenario, cost=1500.0, qalys=104.0, infections=5),
        }
    scenarios["scenario_c_oral_expansion_benchmark"] = {
        "oral_prep": _result(
            "oral_prep",
            "scenario_c_oral_expansion_benchmark",
            cost=1200.0,
            qalys=102.0,
            infections=8,
        ),
        "oral_prep_expansion": _result(
            "oral_prep_expansion",
            "scenario_c_oral_expansion_benchmark",
            cost=1300.0,
            qalys=103.0,
            infections=7,
        ),
    }
    return scenarios


def _write_split_report(tmp_path):
    output_path = tmp_path / "split_hta_report.xlsx"
    baseline_path = tmp_path / "baseline.csv"
    with baseline_path.open("w", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "year",
                "total_population",
                "sexually_active_population",
                "new_infections",
                "prevalence",
                "incidence_rate",
                "hiv_deaths",
                "prep_coverage_Oral_PrEP",
                "prep_coverage_Lenacapavir",
            ],
        )
        writer.writeheader()
        for year, infections in ((1975, 10), (1980, 12), (1985, 14)):
            for q in (0.0, 0.25, 0.5, 0.75):
                writer.writerow({
                    "year": year + q,
                    "total_population": 1000 + year,
                    "sexually_active_population": 800 + year,
                    "new_infections": infections,
                    "prevalence": 0.01,
                    "incidence_rate": 0.001,
                    "hiv_deaths": 1,
                    "prep_coverage_Oral_PrEP": 0.1,
                    "prep_coverage_Lenacapavir": 0.0,
                })
    HTAReportWriter(
        all_results=_all_results(),
        comparator=ScenarioComparator(wtp_threshold=80000, step_fraction=1.0),
        hta_config=_hta_config(),
        base_config={
            "meta": {"model_name": {"value": "test"}, "version": {"value": "split"}},
            "run_control": {"start_year": {"value": 2025}, "step_size": {"value": 12}},
            "simulation_controls": {"deterministic_mode": {"value": True}},
        },
        baseline_csv_path=str(baseline_path),
        output_path=str(output_path),
    ).write()
    assert os.path.exists(output_path)
    return load_workbook(output_path, data_only=False)


def test_split_runtime_report_uses_reviewer_facing_sheet_set(tmp_path):
    wb = _write_split_report(tmp_path)

    required_sheets = {
        "Executive Summary",
        "Run Manifest",
        "Scenario & Variant Register",
        "CEA Results",
        "Cost & Outcome Drivers",
        "Budget Impact",
        "Price & Threshold Analysis",
        "Assumptions & Sources",
        "Baseline Plausibility",
        "Baseline Plausibility Forecast",
    }

    assert required_sheets.issubset(set(wb.sheetnames))
    assert "Decision Summary" not in wb.sheetnames
    assert "Scenario Map" not in wb.sheetnames
    assert any(
        name.startswith("Appendix") or name.startswith("Appendices")
        for name in wb.sheetnames
    )

    assert wb.sheetnames.index("Baseline Plausibility") < wb.sheetnames.index("CEA Results")
    assert wb.sheetnames.index("Baseline Plausibility Forecast") < wb.sheetnames.index("CEA Results")


def test_split_runtime_report_includes_scenario_c_oral_expansion_appendices(tmp_path):
    wb = _write_split_report(tmp_path)

    for sheet_name in ("Appendix - Annual PrEP", "Appendix - Annual Incidence"):
        assert sheet_name in wb.sheetnames
        ws = wb[sheet_name]
        sheet_text = {
            cell.value
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
        }

        assert "Oral PrEP Expansion" in sheet_text
        assert any(
            "scenario_c_oral_expansion_benchmark" in text
            or "Oral PrEP expansion benchmark" in text
            for text in sheet_text
        )


def test_split_runtime_report_register_is_variant_level(tmp_path):
    wb = _write_split_report(tmp_path)
    ws = wb["Scenario & Variant Register"]

    headers = [ws.cell(row=4, column=col).value for col in range(1, 12)]
    assert headers == [
        "Code",
        "Role",
        "Scenario ID",
        "Family",
        "Variant",
        "Label",
        "Comparator Arm",
        "Intervention Arm",
        "Price Sensitivity",
        "Budget Impact",
        "Source Files",
    ]

    scenario_ids = {
        ws.cell(row=row, column=3).value
        for row in range(5, ws.max_row + 1)
        if ws.cell(row=row, column=3).value
    }
    assert "scenario_a_lena_substitution__switch_50" in scenario_ids
    assert "scenario_c_oral_expansion_benchmark" in scenario_ids


def test_split_runtime_report_has_budget_impact_tab(tmp_path):
    wb = _write_split_report(tmp_path)
    ws = wb["Budget Impact"]

    headers = [ws.cell(row=4, column=col).value for col in range(1, 16)]
    assert "Annual Budget Impact Disc" in headers
    assert "Cumulative Budget Impact Disc" in headers
    assert "Annual Infections Averted" in headers
    assert any(
        ws.cell(row=row, column=4).value == "Lena substitution, 50% switch"
        for row in range(5, ws.max_row + 1)
    )
    assert ws.cell(row=5, column=5).number_format == "0"


def test_split_runtime_report_has_compact_baseline_forecast_tab(tmp_path):
    wb = _write_split_report(tmp_path)
    ws = wb["Baseline Plausibility Forecast"]

    headers = [ws.cell(row=4, column=col).value for col in range(1, 8)]
    assert headers == [
        "Arm",
        "Source Scenario",
        "Year",
        "Total Population",
        "New Infections",
        "HIV Prevalence",
        "HIV Deaths",
    ]
    arms = {
        ws.cell(row=row, column=1).value
        for row in range(5, ws.max_row + 1)
        if ws.cell(row=row, column=1).value
    }
    assert {"Attribution No PrEP", "SoC Oral PrEP (no expansion)"}.issubset(arms)
    assert ws.cell(row=5, column=3).number_format == "0"
    assert ws.cell(row=5, column=6).number_format == "0.000%"


def test_split_runtime_report_year_columns_are_plain_years(tmp_path):
    wb = _write_split_report(tmp_path)

    assert wb["Baseline Plausibility"].cell(row=4, column=1).number_format == "0"
    assert wb["Appendix - Annual PrEP"].cell(row=3, column=3).number_format == "0"
    assert wb["Appendix - Annual Incidence"].cell(row=3, column=3).number_format == "0"


def test_split_runtime_report_price_sheet_is_table_first_and_color_coded(tmp_path):
    wb = _write_split_report(tmp_path)
    ws = wb["Price & Threshold Analysis"]

    assert ws["A1"].value == "Price And Threshold Analysis"
    assert not ws._charts
    assert ws.column_dimensions["K"].width >= 24

    detail_header_row = None
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=12).value == "Cost-Effective":
            detail_header_row = row
            break
    assert detail_header_row is not None
    yes_row = detail_header_row + 1
    assert ws.cell(row=yes_row, column=12).value == "Yes"
    assert ws.cell(row=yes_row, column=12).fill.fgColor.rgb in {"00E2EFDA", "E2EFDA"}


def test_split_runtime_report_manifest_formats_only_money_as_money(tmp_path):
    wb = _write_split_report(tmp_path)
    ws = wb["Run Manifest"]

    rows_by_item = {
        ws.cell(row=row, column=2).value: row
        for row in range(5, 17)
        if ws.cell(row=row, column=2).value
    }
    assert ws.cell(row=rows_by_item["WTP threshold"], column=3).number_format == "€#,##0"
    assert ws.cell(row=rows_by_item["Start year"], column=3).number_format != "€#,##0"
    assert ws.cell(row=rows_by_item["Step size"], column=3).number_format != "€#,##0"


def test_split_runtime_report_cea_has_interpretation_key(tmp_path):
    wb = _write_split_report(tmp_path)
    ws = wb["CEA Results"]
    sheet_text = {
        cell.value
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    }

    assert "Interpretation Key" in sheet_text
    assert "Higher cost and lower QALYs than comparator." in sheet_text


def test_split_runtime_report_assumptions_expand_lena_isr_inputs(tmp_path):
    wb = _write_split_report(tmp_path)
    ws = wb["Assumptions & Sources"]
    params = {
        ws.cell(row=row, column=2).value: ws.cell(row=row, column=3).value
        for row in range(5, ws.max_row + 1)
    }

    assert params["lenacapavir_isr_ae.prevalence"] == 0.405
    assert params["lenacapavir_isr_ae.utility_decrement"] == -0.01
    assert params["lenacapavir_isr_ae.cost"] == 0


def test_split_runtime_report_removes_executive_chart_and_retargets_baseline_charts(tmp_path):
    wb = _write_split_report(tmp_path)

    assert not wb["Executive Summary"]._charts
    baseline = wb["Baseline Plausibility"]
    for chart in baseline._charts:
        for series in chart.series:
            assert "Baseline Plausibility" in series.val.numRef.f
            assert "Baseline Key Years" not in series.val.numRef.f


def test_split_runtime_report_drops_legacy_key_figures_and_repeated_cea_blocks(tmp_path):
    wb = _write_split_report(tmp_path)

    assert "Key Figures" not in wb.sheetnames
    assert "CEA Results Summary" not in wb.sheetnames

    workbook_text = {
        cell.value
        for sheet in wb.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    }
    assert "Key Figures — HTA Decision Summary" not in workbook_text
    assert "Backing Data for Key Figures" not in workbook_text


def test_split_report_labels_cost_saving_not_ce_tradeoff(tmp_path):
    output_path = tmp_path / "split_hta_report.xlsx"
    writer = HTAReportWriter(
        all_results=_all_results(),
        comparator=ScenarioComparator(wtp_threshold=80000, step_fraction=1.0),
        hta_config=_hta_config(),
        base_config={
            "meta": {"model_name": {"value": "test"}, "version": {"value": "split"}},
            "run_control": {"start_year": {"value": 2025}, "step_size": {"value": 12}},
            "simulation_controls": {"deterministic_mode": {"value": True}},
        },
        baseline_csv_path="",
        output_path=str(output_path),
    )

    comp = {
        "delta_cost": -1000.0,
        "delta_qaly": -0.02,
        "nmb": -600.0,
        "dominance": "trade-off",
    }

    assert writer._ce_status_from_comparison(comp) == "Cost-saving trade-off; not CE"
    assert not writer._is_ce_status("Cost-saving trade-off; not CE")
