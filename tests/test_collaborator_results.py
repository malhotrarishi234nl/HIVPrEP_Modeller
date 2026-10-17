"""Unit tests for the read-only collaborator result importer; no simulations."""

from __future__ import annotations

from hashlib import sha256
import json

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from collaborator_workbench.results import (
    ResultImportError,
    annualize_model_output,
    build_report,
    inspect_output,
)
from workbench_ui.presentation import (
    display_classification,
    display_model,
    display_probability,
    display_status,
    display_warning,
    implausibility_cue,
)


def _quarterly(seed_shift=0):
    return pd.DataFrame({
        "step": [0, 1, 2, 3],
        "year": [2008.0, 2008.25, 2008.5, 2008.75],
        "new_infections": [1 + seed_shift, 2, 3, 4],
        "adult_hiv_pos": [100, 101, 102, 103 + seed_shift],
        "plhiv_child_0_14": [0, 0, 0, 0],
        "prevalence": [.01, .011, .012, .013],
        "sexually_active_population": [10_000] * 4,
        "incidence_rate": [.1, .2, .3, .4],
        "mystery_value": [9, 8, 7, 6],
    })


def _write_native_run(root, name, model, p, seed, shift=0):
    run_dir = root / name
    run_dir.mkdir()
    csv_path = run_dir / "research_not_production_eligible_run_001.csv"
    _quarterly(shift).to_csv(csv_path, index=False)
    digest = sha256(csv_path.read_bytes()).hexdigest()
    manifest = {
        "schema": "model_run_reproducibility/1.1.0",
        "status": "COMPLETE",
        "classification": "RESEARCH_NOT_PRODUCTION_ELIGIBLE",
        "production_eligible": False,
        "route": {"model": model, "p": p, "seed": seed},
        "scientific_output": {"relative_path": csv_path.name, "sha256": digest},
    }
    (run_dir / "run_reproducibility_manifest.json").write_text(json.dumps(manifest))
    return csv_path


def test_annualization_uses_explicit_flow_stock_and_rate_rules():
    annual, catalog, warnings = annualize_model_output(_quarterly())
    assert annual.to_dict("records") == [{
        "year": 2008,
        "new_infections": 10.0,
        "adult_hiv_pos": 103.0,
        "plhiv_child_0_14": 0.0,
        "prevalence": .013,
        "sexually_active_population": 10_000.0,
        "incidence_rate": .25,
        "PLHIV": 103.0,
    }]
    mystery = catalog.set_index("metric").loc["mystery_value"]
    assert mystery["annual_aggregation"] == "not_annualized"
    assert mystery["included_in_normalized_annual"] == np.bool_(False)
    assert any("mystery_value" in warning for warning in warnings)


def test_incomplete_annual_flow_is_nan_not_partial_total():
    frame = _quarterly()
    frame.loc[2, "new_infections"] = np.nan
    annual, _, _ = annualize_model_output(frame)
    assert np.isnan(annual.loc[0, "new_infections"])


def test_directory_import_verifies_manifests_and_summarizes_routes(tmp_path):
    _write_native_run(tmp_path, "m0-seed1", "M0", None, 1, 0)
    changed = _write_native_run(tmp_path, "m0-seed2", "M0", None, 2, 2)
    _write_native_run(tmp_path, "m1-seed1", "M1", .2, 1, 10)
    bundle = inspect_output(tmp_path)
    assert bundle.kind == "model"
    assert len(bundle.runs) == 3
    assert set(bundle.across_run_summary["route_id"]) == {"M0|p=none", "M1|p=0.2"}
    m0 = bundle.across_run_summary.query("route_id == 'M0|p=none' and metric == 'new_infections'").iloc[0]
    assert m0["n"] == 2
    assert m0["median"] == pytest.approx(11.0)

    # A copied/tampered native result remains inspectable but the report makes
    # the failed provenance binding prominent.
    with changed.open("a") as handle:
        handle.write("\n")
    tampered = inspect_output(tmp_path)
    assert any("does not match" in warning for run in tampered.runs for warning in run.warnings)


def test_generic_csv_is_inspectable_without_guessed_annualization(tmp_path):
    path = tmp_path / "calibration_candidates.csv"
    pd.DataFrame({"candidate": [1, 2], "score": [.4, .3]}).to_csv(path, index=False)
    bundle = inspect_output(path)
    assert bundle.kind == "generic"
    assert bundle.runs == []
    assert bundle.tables[0].frame["score"].tolist() == [.4, .3]


def test_retained_calibration_candidate_imports_full_per_seed_annual_outputs(tmp_path):
    directory = tmp_path / "historical_2008"
    directory.mkdir()
    candidate = {
        "candidate_id": 1787,
        "target_order": ["new_infections:2008", "PLHIV:2008"],
        "target_matrix": [[10, 103], [12, 105]],
        "means": [11, 104],
        "sample_variances": [2, 2],
        "params": {"example": 1},
        "per_seed": [
            {"seed": 41, "annual_outputs": {"new_infections": {"2008": 10, "2009": 9},
                                              "PLHIV": {"2008": 103, "2009": 106}}},
            {"seed": 42, "annual_outputs": {"new_infections": {"2008": 12, "2009": 8},
                                              "PLHIV": {"2008": 105, "2009": 108}}},
        ],
    }
    (directory / "candidate_1787.json").write_text(json.dumps(candidate))
    (directory / "confirmation_summary.csv").write_text("candidate,fit\n1787,good\n")
    (directory / "provenance.json").write_text(json.dumps({
        "classification": "HISTORICAL_RETAINED",
        "scientific_status": "CONFIRMED_GOOD_FIT_FOR_RECORDED_LINEAGE",
        "current_code_compatibility": "NOT_RERUNNABLE_BY_CURRENT_RUN_MODEL",
        "lineage": "2026-09-29",
    }))
    bundle = inspect_output(directory)
    assert len(bundle.runs) == 2
    assert bundle.tables[0].role == "calibration_candidate_annual"
    assert bundle.runs[0].classification == "HISTORICAL_RETAINED"
    assert bundle.runs[0].annual["new_infections"].tolist() == [10.0, 9.0]
    assert set(bundle.annual_long["seed"]) == {41, 42}
    assert any("NOT_RERUNNABLE" in warning for warning in bundle.runs[0].warnings)


def test_report_contains_normalized_data_figures_and_provenance(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    _write_native_run(source, "m1-seed42", "M1", .25, 42)
    bundle = inspect_output(source)
    artifacts = build_report(bundle, tmp_path / "workbench_report", figure_metrics=["new_infections"])
    assert artifacts.index_html.is_file()
    assert artifacts.workbook_xlsx.is_file()
    assert artifacts.annual_csv.is_file()
    assert artifacts.summary_csv.is_file()
    assert artifacts.provenance_json.is_file()
    assert len(artifacts.figures) == 1 and artifacts.figures[0].suffix == ".svg"
    html = artifacts.index_html.read_text()
    assert "M1-p0.25-seed42" in html
    assert "not confidence intervals" in html
    assert "Model in brief" in html
    assert "Disease progression and care" in html
    assert "How the population state is represented" in html
    assert "From fitted history to policy results" in html
    assert "Author's schematic of implemented model logic" in html
    assert "History matching" in html
    assert "Calibration and validation are separate" in html
    assert "Interface treatment" in html
    assert "Adds symptom-triggered diagnosis (M1)" in html
    assert "25%" in html
    assert "What these results show" in html
    assert "How to read the variation" in html
    assert "Downloads and provenance" in html
    assert "Result reference" in html
    assert "source_inventory.csv" in html
    provenance = json.loads(artifacts.provenance_json.read_text())
    assert provenance["runs"][0]["seed"] == 42
    assert provenance["schema"] == "collaborator_results_report/1.1.0"
    assert provenance["result_reference"].startswith("HIVPREP-")
    assert provenance["result_reference"] in html
    workbook = load_workbook(artifacts.workbook_xlsx, read_only=False, data_only=False)
    assert workbook.sheetnames[:5] == [
        "Overview", "Annual results", "Individual runs", "Configuration", "Provenance"
    ]
    assert workbook["Overview"]["A2"].value == "HIV PrEP model results"
    assert len(workbook["Overview"]._images) == 1


def test_machine_metadata_has_plain_labels_without_changing_raw_provenance(tmp_path):
    assert display_status("COMPLETE") == "Run complete"
    assert display_status(
        "VALID_FOR_RECORDED_2026_09_29_LINEAGE_ONLY_NOT_CURRENT_OPTION_B_EVIDENCE"
    ) == "Recorded historical result"
    assert display_classification("HISTORICAL_RETAINED") == "Historical result"
    assert display_model(None) == "Not recorded"
    assert display_probability(None, model="M0") == "Not applicable"
    assert display_warning(
        "Current-code compatibility: INCOMPATIBLE_WITH_CURRENT_OPTION_B_PRODUCTION_ROUTE"
    ) == "This result cannot be rerun by the current model version."

    source = tmp_path / "source"
    source.mkdir()
    _write_native_run(source, "m0-seed42", "M0", None, 42)
    manifest_path = source / "m0-seed42" / "run_reproducibility_manifest.json"
    manifest = json.loads(manifest_path.read_text())
    raw_status = "VALID_FOR_RECORDED_2026_09_29_LINEAGE_ONLY_NOT_CURRENT_OPTION_B_EVIDENCE"
    manifest["status"] = raw_status
    manifest_path.write_text(json.dumps(manifest))

    artifacts = build_report(inspect_output(source), tmp_path / "report")
    html = artifacts.index_html.read_text()
    provenance = json.loads(artifacts.provenance_json.read_text())
    assert "Option B" not in html
    assert "OPTION_B" not in html
    assert "Recorded historical result" in html
    assert provenance["runs"][0]["status"] == raw_status


def test_implausibility_cue_uses_text_and_only_the_documented_threshold():
    within = implausibility_cue(3)
    outside = implausibility_cue(3.01)
    missing = implausibility_cue(None)
    assert (within.label, within.tone) == ("Within the primary threshold", "within")
    assert "I-max <= 3" in within.explanation
    assert (outside.label, outside.tone) == ("Outside the primary threshold", "outside")
    assert missing.label == "Not available"


def test_report_refuses_to_overwrite_nonempty_directory(tmp_path):
    csv_path = tmp_path / "out.csv"
    _quarterly().to_csv(csv_path, index=False)
    bundle = inspect_output(csv_path)
    destination = tmp_path / "report"
    destination.mkdir()
    (destination / "keep.txt").write_text("keep")
    with pytest.raises(ResultImportError, match="not empty"):
        build_report(bundle, destination)
