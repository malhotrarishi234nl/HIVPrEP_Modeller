"""Tests for the Phase 4 evidence package: tables, workbook and verification.

Persisted cover for the 2026-09-23 review findings:

* blank source cells must never render as 0 in the workbook (a blank on a price
  column reads as "a price of zero", the inversion of "no price exists");
* the formula verifier must actually catch wrong text, wrong numbers and dead
  references, and must NOT accept 0 for a blank;
* the structural scan is a static check, not a render test, and says so.

These run no simulation. They rebuild the workbook into tmp_path.
"""

from __future__ import annotations

import math
from pathlib import Path
import shutil

import pandas as pd
import pytest

import tools.build_abstract_evidence_phase4 as phase4

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "outputs/abstract_evidence_20260923"

pytestmark = pytest.mark.skipif(
    not (PACKAGE / "headline_results.csv").exists(),
    reason="Phase 4 tables not built; run tools/build_abstract_evidence_phase4.py",
)


@pytest.fixture(scope="module")
def tables() -> dict[str, pd.DataFrame]:
    return {
        name: pd.read_csv(PACKAGE / name, float_precision="round_trip")
        for name in phase4.DATA_SHEETS.values()
    }


@pytest.fixture(scope="module")
def built(tables, tmp_path_factory):
    out = tmp_path_factory.mktemp("phase4")
    path, expectations = phase4.build_workbook(out, tables)
    return path, expectations


@pytest.fixture(scope="module")
def checks(built) -> pd.DataFrame:
    path, expectations = built
    return phase4.verify_workbook(path, expectations)


# ---------------------------------------------------------------------------
# The blank-to-zero defect
# ---------------------------------------------------------------------------
def test_blank_source_cells_are_guarded_not_referenced_bare(built):
    """A bare reference to a blank cell renders as 0. Every blank must be wrapped."""
    _, expectations = built
    guarded = [e for e in expectations if e.get("blank_guarded")]
    assert guarded, "expected some blank-capable cells in the headline tab"
    for item in guarded:
        assert item["formula"].startswith("=IF(ISBLANK("), item
        assert item["expected"] in (
            phase4.NO_PRICE_SENTINEL, phase4.NOT_APPLICABLE_SENTINEL
        )


def test_no_price_scenarios_do_not_render_as_zero(checks):
    """The three legacy add-on rows must read as text, never as a price of 0."""
    guarded = checks[checks.blank_guarded]
    assert len(guarded) > 0
    assert not (guarded.observed.astype(str) == "0").any()
    prices = guarded[guarded.expected == phase4.NO_PRICE_SENTINEL]
    assert len(prices) == 6, "price value and reduction, three legacy add-on scenarios"
    assert (prices.observed == phase4.NO_PRICE_SENTINEL).all()


def test_price_blank_sentinels_follow_row_semantics(checks):
    """No-root and not-applicable rows must never share a display sentinel."""
    headline = pd.read_csv(PACKAGE / "headline_results.csv", float_precision="round_trip")
    priceish = {
        column for column in headline.columns
        if ("price" in column and "quantity" not in column and "stability" not in column
            and "exists" not in column and "base_annual" not in column)
    }
    assert priceish <= phase4.PRICE_COLUMNS, sorted(priceish - phase4.PRICE_COLUMNS)
    checked = checks.set_index(["sheet", "cell"])
    blank_rows = 0
    for index, row in headline.iterrows():
        if not pd.isna(row.pooled_nmb_zero_price_80k):
            continue
        blank_rows += 1
        expected = {
            "no_non_negative_root": phase4.NO_PRICE_SENTINEL,
            "not_applicable": phase4.NOT_APPLICABLE_SENTINEL,
        }[row.pooled_price_quantity_name_80k]
        for column in ("V", "W"):
            check = checked.loc[("Headline results", f"{column}{index + 4}")]
            assert check.expected == expected
            assert check.observed == expected
    assert blank_rows == 11


def _corrupt_formula(built, tmp_path, item, formula):
    """Change a real formula in a disposable workbook copy."""
    from openpyxl import load_workbook

    source, expectations = built
    corrupted = tmp_path / "corrupted.xlsx"
    shutil.copyfile(source, corrupted)
    workbook = load_workbook(corrupted)
    workbook[item["sheet"]][item["cell"]] = formula
    workbook.save(corrupted)
    return phase4.verify_workbook(corrupted, expectations)


# ---------------------------------------------------------------------------
# The verifier must have teeth
# ---------------------------------------------------------------------------
def test_verifier_rejects_zero_for_a_blank(built, tmp_path):
    """Regression: the old verifier accepted 0 as equivalent to blank."""
    _, expectations = built
    target = next(e for e in expectations if e.get("blank_guarded"))
    result = _corrupt_formula(built, tmp_path, target, "=0")
    failed = result[~result.passes]
    assert len(failed) == 1
    assert failed.iloc[0]["cell"] == target["cell"]


@pytest.mark.parametrize(
    "formula,label",
    [('="THIS SHOULD NOT MATCH"', "wrong text"),
     ("=#REF!", "dead reference")],
)
def test_verifier_catches_corruption(built, tmp_path, formula, label):
    _, expectations = built
    target = next(e for e in expectations
                  if isinstance(e["expected"], str) and not e.get("blank_guarded"))
    result = _corrupt_formula(built, tmp_path, target, formula)
    assert int((~result.passes).sum()) == 1, label
    assert result.loc[~result.passes, "cell"].iloc[0] == target["cell"]


def test_verifier_catches_a_wrong_number(built, tmp_path):
    _, expectations = built
    target = next(e for e in expectations
                  if isinstance(e["expected"], float)
                  and not math.isnan(e["expected"]) and e["expected"] != 0)
    result = _corrupt_formula(built, tmp_path, target, "=12345")
    assert int((~result.passes).sum()) == 1
    assert result.loc[~result.passes, "cell"].iloc[0] == target["cell"]


def test_all_formulas_verify_on_the_real_workbook(checks):
    failures = checks[~checks.passes]
    assert failures.empty, failures.head().to_string()
    assert len(checks) > 2000


# ---------------------------------------------------------------------------
# Workbook mechanics
# ---------------------------------------------------------------------------
def test_workbook_forces_recalculation_on_load(built):
    """No cached results are stored, so viewers must be told to calculate."""
    from openpyxl import load_workbook

    path, _ = built
    assert load_workbook(path).calculation.fullCalcOnLoad is True


def test_structural_scan_is_named_and_documented_as_static(built):
    """Regression: the old name claimed a render test it did not perform."""
    path, expectations = built
    assert not hasattr(phase4, "scan_workbook"), "old misleading name must be gone"
    doc = phase4.scan_workbook_structure.__doc__ or ""
    assert "NOT a render test" in doc
    scan = phase4.scan_workbook_structure(path, expectations)
    assert "populated" in scan.columns and "renders" not in scan.columns
    assert scan.passes.all()
    assert int(scan.excel_error_cells.sum()) == 0


def test_every_presentation_tab_is_formula_linked(built):
    path, expectations = built
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    tabs = [s for s in workbook.sheetnames if not s.startswith("Data_")]
    assert len(tabs) == 6
    linked = {e["sheet"] for e in expectations}
    assert set(tabs) <= linked


# ---------------------------------------------------------------------------
# Memo scope
# ---------------------------------------------------------------------------
def test_memo_scopes_the_no_break_even_claim_to_convention_and_threshold():
    """Regression: the memo stated this without naming convention or threshold."""
    memo = (PACKAGE / "ABSTRACT_INPUTS.md").read_text(encoding="utf-8")
    prices = pd.read_csv(PACKAGE / "price_thresholds.csv", float_precision="round_trip")
    legacy_addon = prices[
        (prices.structure == "legacy")
        & prices.deployment_family_label.str.contains("add-on")
    ]
    # The data: no root anywhere under the configured convention; roots exist at
    # 50k and 80k under no-ISR.
    configured = legacy_addon[legacy_addon.ae_convention == "configured_isr_nodule"]
    assert not configured.root_exists.any()
    no_isr = legacy_addon[legacy_addon.ae_convention == "symmetric_no_isr"]
    assert not no_isr[no_isr.wtp == 20000.0].root_exists.any()
    assert no_isr[no_isr.wtp.isin([50000.0, 80000.0])].root_exists.all()
    # The memo must carry that scope.
    assert "configured adverse-event convention" in memo
    assert "symmetric no-ISR convention that remains" in memo


def test_memo_quotes_only_numbers_that_exist_in_the_tables():
    memo = (PACKAGE / "ABSTRACT_INPUTS.md").read_text(encoding="utf-8")
    headline = pd.read_csv(PACKAGE / "headline_results.csv", float_precision="round_trip")
    configured = headline[headline.ae_convention == "configured_isr_nodule"]

    def value(structure, family, level, column):
        row = configured[
            (configured.structure == structure)
            & (configured.deployment_family == family)
            & (configured.deployment_level == level)
        ]
        return float(row.iloc[0][column])

    assert f"{value('legacy', 'lena_substitution', 100, 'mean_delta_cost') / 1e6:,.1f}" in memo
    assert f"{value('legacy', 'lena_add_on', 3000, 'mean_infections_averted'):.1f}" in memo
