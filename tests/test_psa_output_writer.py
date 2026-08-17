"""
tests/test_psa_output_writer.py — V11.11.1

Unit tests for PSA output writer: summary statistics, rank correlations,
convergence calculations, and Excel structure.

5 tests:
  - test_outcome_summary_statistics
  - test_rank_correlation_known_monotonic
  - test_rank_correlation_independent
  - test_convergence_data_shape
  - test_excel_workbook_structure
"""

import os
import sys
import tempfile

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from psa_output_writer import PSAOutputWriter


# ── Mock helpers ─────────────────────────────────────────────────────────────


class MockSampler:
    def get_parameter_summary(self):
        return pd.DataFrame({
            "path": ["a.b", "c.d"],
            "tier": ["uncertain", "uncertain"],
            "value": [0.5, 0.1],
            "bounds_lo": [0.0, 0.0],
            "bounds_hi": [1.0, 0.5],
            "dist_type": ["Beta", "Beta"],
            "dist_params": ["{'alpha': 5, 'beta': 5}", "{'alpha': 1, 'beta': 9}"],
            "source": ["test", "test"],
        })


def _minimal_psa_config(output_dir):
    """Minimal psa_config that disables all plots to keep tests fast."""
    return {
        "outputs": {
            "raw_draws": {"enabled": True},
            "summary_excel": {"enabled": True},
            "plots": {
                "format": {"value": "png"},
                "dpi": {"value": 72},
                "tornado": {"enabled": False, "metric": "metric_a", "top_n": 5},
                "histogram": {"enabled": False, "metrics": []},
                "convergence": {"enabled": False, "metric": "metric_a"},
            },
            "directory": {"value": str(output_dir)},
        }
    }


def _make_writer(draws_df, outcomes_df, output_dir):
    psa_config = _minimal_psa_config(output_dir)
    return PSAOutputWriter(
        draws_df=draws_df,
        outcomes_df=outcomes_df,
        sampler=MockSampler(),
        psa_config=psa_config,
        output_dir=output_dir,
    )


# ── Tests ─────────────────────────────────────────────────────────────────────


def test_outcome_summary_statistics():
    """
    _compute_outcome_summary() returns correct summary stats for a known distribution.
    """
    rng = np.random.default_rng(42)
    vals_a = rng.normal(100, 10, 100)
    vals_b = rng.normal(50, 5, 100)

    outcomes_df = pd.DataFrame({
        "metric_a": vals_a,
        "metric_b": vals_b,
    })
    draws_df = pd.DataFrame({"param_x": rng.random(100)})

    with tempfile.TemporaryDirectory() as out_dir:
        writer = _make_writer(draws_df, outcomes_df, out_dir)
        summary = writer._compute_outcome_summary()

    assert isinstance(summary, pd.DataFrame)
    assert "metric" in summary.columns
    row_a = summary[summary["metric"] == "metric_a"].iloc[0]

    # Mean within 5% of true mean
    assert abs(row_a["mean"] - np.mean(vals_a)) / np.mean(vals_a) < 0.05
    # Std within 10% of true std
    assert abs(row_a["std"] - np.std(vals_a, ddof=1)) / np.std(vals_a, ddof=1) < 0.10
    # p2.5 <= mean <= p97.5
    assert row_a["p2.5"] <= row_a["mean"] <= row_a["p97.5"]
    # Required columns present
    required = {"metric", "mean", "median", "std", "p2.5", "p97.5", "min", "max"}
    assert required.issubset(set(summary.columns))


def test_rank_correlation_known_monotonic():
    """
    Perfectly monotonic relationship → Spearman rho = 1.0.
    """
    n = 50
    param_x = np.arange(1, n + 1, dtype=float)
    outcome_y = 2.0 * param_x  # perfectly linear (monotonic)

    draws_df = pd.DataFrame({"param_x": param_x})
    outcomes_df = pd.DataFrame({"outcome_y": outcome_y})

    with tempfile.TemporaryDirectory() as out_dir:
        writer = _make_writer(draws_df, outcomes_df, out_dir)
        corr_df = writer._compute_rank_correlations("outcome_y")

    row = corr_df[corr_df["parameter"] == "param_x"].iloc[0]
    assert abs(row["spearman_rho"] - 1.0) < 1e-9


def test_rank_correlation_independent():
    """
    Independent random variables → |Spearman rho| should be near 0.
    """
    rng_a = np.random.default_rng(42)
    rng_b = np.random.default_rng(99)
    n = 200

    draws_df = pd.DataFrame({"param_x": rng_a.random(n)})
    outcomes_df = pd.DataFrame({"outcome_y": rng_b.random(n)})

    with tempfile.TemporaryDirectory() as out_dir:
        writer = _make_writer(draws_df, outcomes_df, out_dir)
        corr_df = writer._compute_rank_correlations("outcome_y")

    row = corr_df[corr_df["parameter"] == "param_x"].iloc[0]
    assert abs(row["spearman_rho"]) < 0.2


def test_convergence_data_shape():
    """
    _compute_convergence_data returns correct columns and final n_draws == N.
    ci_lo <= running_mean <= ci_hi for all rows with > 1 draw.
    """
    rng = np.random.default_rng(7)
    n = 100
    outcomes_df = pd.DataFrame({"metric_a": rng.normal(50, 10, n)})
    draws_df = pd.DataFrame({"param_x": rng.random(n)})

    with tempfile.TemporaryDirectory() as out_dir:
        writer = _make_writer(draws_df, outcomes_df, out_dir)
        conv = writer._compute_convergence_data("metric_a")

    assert set(conv.columns) >= {"n_draws", "running_mean", "ci_lo", "ci_hi"}
    assert conv["n_draws"].iloc[-1] == n

    # For all rows where CI is defined (n > 1), ci_lo <= mean <= ci_hi
    for _, row in conv.iterrows():
        if not (np.isnan(row["ci_lo"]) or np.isnan(row["ci_hi"])):
            assert row["ci_lo"] <= row["running_mean"] <= row["ci_hi"], (
                f"CI violated at n_draws={row['n_draws']}: "
                f"[{row['ci_lo']:.4f}, {row['ci_hi']:.4f}] mean={row['running_mean']:.4f}"
            )


def test_excel_workbook_structure():
    """
    write_all() produces an Excel workbook with required sheets and correct row count.
    """
    import openpyxl

    rng = np.random.default_rng(42)
    n = 5
    draws_df = pd.DataFrame({
        "param_a": rng.random(n),
        "param_b": rng.random(n),
    })
    outcomes_df = pd.DataFrame({
        "metric_x": rng.random(n),
        "metric_y": rng.random(n),
    })

    with tempfile.TemporaryDirectory() as out_dir:
        writer = _make_writer(draws_df, outcomes_df, out_dir)
        writer.write_all()

        wb = openpyxl.load_workbook(os.path.join(out_dir, "psa_summary.xlsx"))
        sheet_names = wb.sheetnames

    assert "parameter_summary" in sheet_names
    assert "outcome_summary" in sheet_names
    assert "raw_draws" in sheet_names
    assert "raw_outcomes" in sheet_names

    # raw_draws has 5 data rows (1 header + 5 data rows = 6 total rows in sheet)
    ws = wb["raw_draws"]
    data_rows = ws.max_row - 1  # subtract header
    assert data_rows == n
