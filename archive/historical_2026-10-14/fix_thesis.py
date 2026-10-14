#!/usr/bin/env python3
"""
Fix THESIS_FINAL.docx:
  TASK 1: Fix List of Tables and Figures
  TASK 2: Highlight orphan references in yellow
"""

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from copy import deepcopy
import re

DOC_PATH = '/sessions/nifty-vigilant-mayer/mnt/Modeller/v11/THESIS_FINAL.docx'

doc = Document(DOC_PATH)

changes_made = []

# ============================================================
# TASK 1: Fix the List of Tables and Figures
# ============================================================

print("=" * 60)
print("TASK 1: Fixing List of Tables and Figures")
print("=" * 60)

# Helper: find paragraph by text match in a range
def find_para(start, end, substring):
    for i in range(start, min(end, len(doc.paragraphs))):
        if substring in doc.paragraphs[i].text:
            return i
    return None

# Helper: clear all runs in a paragraph and set new text (preserving style)
def set_paragraph_text(para_idx, new_text):
    p = doc.paragraphs[para_idx]
    old_text = p.text
    # Clear existing runs
    for run in p.runs:
        run.text = ''
    # Set the first run's text (or create one if needed)
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)
    return old_text

# Helper: insert a paragraph after a given paragraph index
def insert_paragraph_after(para_idx, text, style_name='Normal'):
    """Insert a new paragraph after the given paragraph index."""
    p = doc.paragraphs[para_idx]
    # Create new paragraph element
    new_p = deepcopy(p._element)
    # Clear the new paragraph's text
    from docx.oxml.ns import qn
    for r in new_p.findall(qn('w:r')):
        new_p.remove(r)
    # Add run with text
    from lxml import etree
    r_elem = new_p.makeelement(qn('w:r'), {})
    t_elem = r_elem.makeelement(qn('w:t'), {})
    t_elem.text = text
    # Preserve spaces
    t_elem.set(qn('xml:space'), 'preserve')
    r_elem.append(t_elem)
    new_p.append(r_elem)
    # Insert after the current paragraph
    p._element.addnext(new_p)
    return new_p

# --- Fix 1: Table 5 - add period after "weights" ---
idx = find_para(30, 55, 'Table 5. QALY weights')
if idx is not None:
    p = doc.paragraphs[idx]
    # The first run has "Table 5. QALY weights", add period
    for run in p.runs:
        if 'QALY weights' in run.text and not run.text.strip().endswith('.'):
            old = run.text
            run.text = run.text.replace('QALY weights', 'QALY weights.')
            msg = f"  P{idx}: Added period: '{old}' -> '{run.text}'"
            print(msg)
            changes_made.append(msg)
            break

# --- Fix 2: Table 8 - expand caption ---
idx = find_para(30, 55, 'Table 8. Price sensitivity analysis: Scenario A 10% switch')
if idx is not None:
    p = doc.paragraphs[idx]
    for run in p.runs:
        if 'Table 8. Price sensitivity analysis' in run.text:
            old = run.text
            run.text = 'Table 8. Price sensitivity analysis: Scenario A 10% switch, Lenacapavir injection cost (base = €41,110/year)'
            msg = f"  P{idx}: Updated Table 8 caption"
            print(msg)
            changes_made.append(msg)
            break

# --- Fix 3: Table 9 - expand caption ---
idx = find_para(30, 55, 'Table 9. Cumulative discounted budget impact')
if idx is not None:
    p = doc.paragraphs[idx]
    for run in p.runs:
        if 'Table 9. Cumulative discounted budget impact' in run.text:
            old = run.text
            run.text = 'Table 9. Cumulative discounted budget impact: Scenario A 10% switch vs oral PrEP SOC (EUR, 4%/year discount)'
            msg = f"  P{idx}: Updated Table 9 caption"
            print(msg)
            changes_made.append(msg)
            break

# --- Fix 4: Figure 5 - change "PSA" to "probabilistic sensitivity analysis" ---
idx = find_para(30, 55, 'Figure 5. Cost-effectiveness plane')
if idx is not None:
    p = doc.paragraphs[idx]
    # Runs are: 'Figure 5. Cost-effectiveness plane for ', 'Lenacapavir', ' versus oral PrEP in the ', 'PSA', '.', '\t', '30'
    for run in p.runs:
        if run.text == 'PSA':
            run.text = 'probabilistic sensitivity analysis'
            msg = f"  P{idx}: Changed 'PSA' to 'probabilistic sensitivity analysis'"
            print(msg)
            changes_made.append(msg)
            break

# --- Fix 5: Appendix Figure E.1 - fix caption ---
# Current: "PR sensitivity-analysis output for the selected scenario set."
# Target: "Probabilistic sensitivity-analysis output for Oral PrEP attribution."
idx = find_para(30, 55, 'Appendix Figure E.1')
if idx is not None:
    p = doc.paragraphs[idx]
    # Runs are: 'Appendix Figure E.1. ', 'PR', ' sensitivity-analysis output for the selected scenario set.', '\t', '50'
    for run in p.runs:
        if run.text == 'PR':
            run.text = 'Probabilistic'
            msg = f"  P{idx}: Changed 'PR' to 'Probabilistic'"
            print(msg)
            changes_made.append(msg)
        if 'sensitivity-analysis output for the selected scenario set.' in run.text:
            run.text = ' sensitivity-analysis output for Oral PrEP attribution.'
            msg = f"  P{idx}: Changed 'for the selected scenario set' to 'for Oral PrEP attribution'"
            print(msg)
            changes_made.append(msg)

# --- Fix 6: Appendix Figure P.1 (second one = Garnett) should be P.3 ---
# P51 currently says "Appendix Figure P.1. Garnett & Anderson Mixing."
# Body says "Appendix Figure P.3. Garnett & Anderson Mixing"
idx = find_para(30, 55, 'Garnett')
if idx is not None:
    p = doc.paragraphs[idx]
    # Fix the ".1." run to ".3."
    for run in p.runs:
        if run.text == '.1. ':
            run.text = '.3. '
            msg = f"  P{idx}: Changed 'P.1' to 'P.3' for Garnett & Anderson Mixing"
            print(msg)
            changes_made.append(msg)
            break

# --- Fix 7: Add Appendix Figure E.2 after E.1 ---
# Need to insert after the E.1 entry
idx_e1 = find_para(30, 55, 'Appendix Figure E.1')
if idx_e1 is not None:
    new_text = 'Appendix Figure E.2. Deterministic sensitivity-analysis output for Oral PrEP attribution.\t50'
    insert_paragraph_after(idx_e1, new_text)
    msg = f"  After P{idx_e1}: Inserted 'Appendix Figure E.2. Deterministic sensitivity-analysis output for Oral PrEP attribution.\\t50'"
    print(msg)
    changes_made.append(msg)

print()
print(f"TASK 1 complete: {len(changes_made)} changes made")

# Verify the list looks correct now
print("\n--- Updated List of Tables and Figures ---")
for i in range(30, 58):
    if i < len(doc.paragraphs):
        text = doc.paragraphs[i].text.strip()
        if text:
            print(f"  P{i}: {text[:120]}")


# ============================================================
# TASK 2: Highlight orphan references in yellow
# ============================================================

print("\n" + "=" * 60)
print("TASK 2: Highlighting orphan references")
print("=" * 60)

# Define orphan references with unique matching substrings
# Each tuple: (display_name, match_substring, exclusion_substring_or_None)
orphan_refs = [
    ("CDC (2014)", "Centers for Disease Control and Prevention. (2014)", None),
    ("Cohen MS (2011)", "Cohen MS, Chen YQ, McCauley M", None),
    ("Coutinho RA (1998)", "Coutinho RA. (1998)", None),
    ("EMA (2025b)", "European Medicines Agency. (2025b)", None),
    ("Farmacotherapeutisch Kompas (2026)", "Farmacotherapeutisch Kompas. (2026)", None),
    ("Hakkaart-van Roijen L (2024)", "Hakkaart-van Roijen L", None),
    ("Hollingsworth TD (2008)", "Hollingsworth TD, Anderson RM, Fraser C. (2008)", None),
    ("Kayaert L, Sarink D (2023)", "Kayaert L, Sarink D", None),
    ("Koopmanschap MA (1995)", "Koopmanschap", None),  # only one Koopmanschap
    ("Kuipers EJ (2023)", "Kuipers EJ. (2023)", None),
    ("Patel P (2014)", "Patel, P., Borkowf", None),
    ("Roberts M, Russell LB (2012)", "Roberts M, Russell LB", None),
    ("Rodger AJ, Cambiano V (2019)", "Rodger AJ, Cambiano V", None),
    # Shaik JS not found in document - skip
    ("Wang H, Zimmermann (2024)", "Wang H, Zimmermann HML", None),
    ("World Bank (2026)", "World Bank. (2026)", None),
    ("ZIN (2024b)", "Zorginstituut Nederland. (2024b)", None),
    ("Nederlandse PrEP-richtlijn (2025)", "Nederlandse multidisciplinaire PrEP-richtlijn werkgroep. (2025)", None),
]

# Find the References section
ref_start = None
ref_end = len(doc.paragraphs)
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'References' and 'Heading' in p.style.name:
        ref_start = i
    # Find end of references (next major heading after references)
    if ref_start is not None and i > ref_start + 5:
        if 'Heading 1' in p.style.name and p.text.strip() != 'References':
            ref_end = i
            break

print(f"References section: P{ref_start} to P{ref_end}")
highlight_count = 0

for display_name, match_substr, exclude_substr in orphan_refs:
    found = False
    for i in range(ref_start, ref_end):
        p = doc.paragraphs[i]
        text = p.text
        if match_substr in text:
            if exclude_substr and exclude_substr in text:
                continue  # skip exclusions
            # Apply yellow highlight to all runs
            for run in p.runs:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            found = True
            highlight_count += 1
            msg = f"  P{i}: Highlighted {display_name}"
            print(msg)
            changes_made.append(msg)
            break
    if not found:
        print(f"  WARNING: Could not find '{display_name}' (match: '{match_substr}')")

# Handle Shaik JS separately since it's not in the document
print("  NOTE: Shaik JS (2020) not found in document - skipping")

# Handle Kayaert - need to be careful not to match El-Jobouri ref which also has Kayaert
# Let me verify
for i in range(ref_start, ref_end):
    p = doc.paragraphs[i]
    if 'Kayaert' in p.text:
        has_highlight = any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in p.runs)
        print(f"  Kayaert check P{i}: highlighted={has_highlight}, starts={repr(p.text[:60])}")

print(f"\nTASK 2 complete: {highlight_count} references highlighted in yellow")

# ============================================================
# Save
# ============================================================
doc.save(DOC_PATH)
print(f"\nDocument saved to {DOC_PATH}")
print(f"Total changes: {len(changes_made)}")
