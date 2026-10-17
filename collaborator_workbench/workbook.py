"""Create the primary Excel review workbook for a result bundle."""

from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
import tempfile
from typing import Any, Iterable, Mapping, Sequence

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "20354A"
BLUE = "2F80ED"
LIGHT_BLUE = "EAF2FB"
LIGHT_GREY = "F2F4F7"
MID_GREY = "D7DCE3"
TEXT = "243042"
GREEN = "E7F3E9"
RED = "FBE9E7"
AMBER = "FFF2CC"
WHITE = "FFFFFF"
THIN_GREY = Side(style="thin", color=MID_GREY)

PARAMETER_LABELS = {
    "population.p_msm.value": "Share of the population in the MSM group",
    "population.epidemic_seed.seed_prevalence.value": "Initial HIV prevalence",
    "population.epidemic_seed.msm_fraction.value": "Initial HIV cases in the MSM group",
    "transmission_model.mixing.epsilon_hetero.value": "Mixing within heterosexual activity groups",
    "transmission_model.mixing.epsilon_msm.value": "Mixing within MSM activity groups",
    "transmission_model.disease.transmission.advanced.beta_per_route.vaginal_mtf.chronic.value": "Transmission beta: vaginal sex, male to female",
    "transmission_model.disease.transmission.advanced.beta_per_route.vaginal_ftm.chronic.value": "Transmission beta: vaginal sex, female to male",
    "transmission_model.disease.transmission.advanced.beta_per_route.anal_receptive.chronic.value": "Transmission beta: anal sex, receptive partner",
    "transmission_model.disease.transmission.advanced.beta_per_route.anal_insertive.chronic.value": "Transmission beta: anal sex, insertive partner",
    "transmission_model.disease.transmission.advanced.stage_multiplier_acute.value": "Acute-stage transmission multiplier",
    "transmission_model.disease.transmission.advanced.stage_multiplier_aids.value": "AIDS-stage transmission multiplier",
    "transmission_model.disease.transmission.advanced.msm_versatility.value": "Share of MSM represented as sexually versatile",
    "transmission_model.disease.advanced.art_mortality_reduction.value": "Mortality reduction while receiving ART",
}

CHART_LABELS = {
    "new_infections": "New infections",
    "new_diagnoses": "New diagnoses",
    "PLHIV": "People living with HIV",
    "on_art_stock": "People receiving ART",
    "prevalence": "HIV prevalence",
}

KEY_METRICS = (
    "PLHIV", "new_infections", "new_diagnoses", "new_diagnoses_msm",
    "aids_deaths", "on_art_stock", "prevalence",
)

AGGREGATION_LABELS = {
    "annual_sum": "Sum of timesteps",
    "year_end": "Year-end value",
    "year_start": "Start-of-year value",
    "annual_mean": "Mean across timesteps",
    "reported_annual": "Reported annual value",
}


def _route_label(value: Any) -> str:
    raw = str(value)
    if raw == "unspecified|p=none":
        return "Recorded fit"
    return raw.replace("|p=none", "").replace("|p=", ", p=")


def _unit_label(metric: Any, unit: Any) -> str:
    if str(unit) != "unspecified":
        return str(unit)
    if str(metric) in {"PLHIV", "new_infections", "new_diagnoses", "new_diagnoses_msm", "aids_deaths", "on_art_stock"}:
        return "people"
    if str(metric) == "prevalence":
        return "proportion"
    return "not recorded"


def _display_path(path: Path, bundle: Any) -> str:
    root = Path(__file__).resolve().parents[1]
    for base in (root, bundle.source_path.parent):
        try:
            return path.resolve().relative_to(base.resolve()).as_posix()
        except ValueError:
            continue
    return path.name


def _plain_label(value: Any) -> str:
    return str(value).replace("_", " ").strip().capitalize()


def _value_at(payload: Mapping[str, Any], path: str) -> Any:
    node: Any = payload
    for part in path.split("."):
        if not isinstance(node, Mapping) or part not in node:
            return None
        node = node[part]
    return node


def _prepare_sheet(ws, *, freeze: str | None = None) -> None:
    ws.sheet_view.showGridLines = False
    if freeze:
        ws.freeze_panes = freeze
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _title(ws, text: str, subtitle: str | None = None) -> int:
    ws["A2"] = text
    ws["A2"].font = Font(name="Arial", size=16, bold=True, color=TEXT)
    ws["A3"].fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[3].height = 3
    if subtitle:
        ws["A5"] = subtitle
        ws["A5"].font = Font(name="Arial", size=10, italic=True, color="5E6878")
        return 7
    return 5


def _write_table(
    ws, start_row: int, headers: Sequence[str], rows: Iterable[Sequence[Any]],
    *, table_name: str, widths: Sequence[float] | None = None,
) -> int:
    rows = list(rows)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, col, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN_GREY)
    for row_index, row in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row_index, col, value)
            cell.font = Font(name="Arial", size=10, color=TEXT)
            cell.alignment = Alignment(
                horizontal="right" if isinstance(value, (int, float)) else "left",
                vertical="center",
            )
            cell.border = Border(bottom=Side(style="hair", color="E6E9EE"))
    end_row = start_row + max(len(rows), 1)
    if rows:
        table = Table(
            displayName=table_name,
            ref=f"A{start_row}:{get_column_letter(len(headers))}{end_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
    return end_row


def _format_number_cells(ws, start_row: int, end_row: int, columns: Sequence[int]) -> None:
    for row in range(start_row, end_row + 1):
        for column in columns:
            ws.cell(row, column).number_format = "#,##0.000"


def _chart_png(bundle: Any, metric: str, destination: Path) -> bool:
    summary = bundle.across_run_summary
    data = summary[summary["metric"] == metric].copy()
    if data.empty:
        return False
    fig, ax = plt.subplots(figsize=(7.2, 3.6), dpi=150)
    for route_id, route in data.groupby("route_id", dropna=False, sort=True):
        route = route.sort_values("year")
        x = route["year"].to_numpy(dtype=float)
        median = route["median"].to_numpy(dtype=float)
        label = _route_label(route_id)
        if int(route["n"].max()) > 1:
            ax.fill_between(
                x, route["q05"].to_numpy(dtype=float),
                route["q95"].to_numpy(dtype=float), alpha=0.15,
            )
        ax.plot(x, median, linewidth=2.2, label=label)
    first = data.iloc[0]
    ax.set_title(
        CHART_LABELS.get(metric, str(first.get("label") or _plain_label(metric))),
        loc="left",
    )
    ax.set_xlabel("Year")
    ax.set_ylabel(str(first.get("unit") or ""))
    ax.grid(axis="y", alpha=0.18)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, loc="upper left")
    fig.tight_layout()
    fig.savefig(destination, bbox_inches="tight")
    plt.close(fig)
    return True


def _configuration_rows(bundle: Any) -> list[list[Any]]:
    root = Path(__file__).resolve().parents[1]
    contract_path = root / "configs" / "collaborator_parameter_contract.json"
    canonical_path = root / "configs" / "V13.22_1980_recalibration_candidate_stochastic_m0_m1_20261015.json"
    try:
        contract = json.loads(contract_path.read_text(encoding="utf-8"))
        canonical = json.loads(canonical_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    effective = None
    for run in bundle.runs:
        if run.manifest_path:
            candidate = run.manifest_path.parent / "effective_config.json"
            if candidate.is_file():
                try:
                    effective = json.loads(candidate.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError):
                    pass
                if effective is not None:
                    break
    if effective is None:
        return []
    rows: list[list[Any]] = []
    skipped: set[str] = set()
    for rule in contract.get("parameters", []):
        path = str(rule["path"])
        if path in skipped:
            continue
        linked = rule.get("must_equal")
        if linked:
            skipped.add(str(linked))
        baseline = _value_at(canonical, path)
        value = _value_at(effective, path)
        group = (
            "Transmission" if "transmission" in path and "art_mortality" not in path
            else "Progression and treatment" if "art_mortality" in path
            else "Mixing and epidemic start"
        )
        rows.append([
            group, PARAMETER_LABELS.get(path, _plain_label(path.split(".")[-2])),
            baseline, value, rule.get("unit"), "Changed" if value != baseline else "Unchanged",
        ])
    return rows


def build_workbook(
    bundle: Any, destination: str | Path, *, reference_id: str,
    generated_utc: str, figure_metrics: Sequence[str] | None = None,
) -> Path:
    """Write one portable workbook whose first sheet is useful without other files."""
    output = Path(destination).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    _prepare_sheet(overview)
    overview.sheet_properties.tabColor = NAVY
    row = _title(
        overview, "HIV PrEP model results",
        "Summary, stochastic variation, exact values, configuration and provenance",
    )

    summary = bundle.across_run_summary
    years = pd.to_numeric(summary.get("year", pd.Series(dtype=float)), errors="coerce").dropna()
    routes = int(summary["route_id"].nunique()) if not summary.empty else 0
    metrics = int(summary["metric"].nunique()) if not summary.empty else 0
    overview.cell(row, 1, "Result reference").font = Font(bold=True, color=TEXT)
    overview.cell(row, 2, reference_id)
    overview.cell(row + 1, 1, "Period").font = Font(bold=True, color=TEXT)
    overview.cell(row + 1, 2, f"{int(years.min())}–{int(years.max())}" if not years.empty else "Not available")
    overview.cell(row + 2, 1, "Model setups").font = Font(bold=True, color=TEXT)
    overview.cell(row + 2, 2, routes)
    overview.cell(row + 3, 1, "Trajectories").font = Font(bold=True, color=TEXT)
    overview.cell(row + 3, 2, len(bundle.runs))
    overview.cell(row + 4, 1, "Outcomes").font = Font(bold=True, color=TEXT)
    overview.cell(row + 4, 2, metrics)
    overview.cell(row + 5, 1, "Generated (UTC)").font = Font(bold=True, color=TEXT)
    try:
        generated_label = datetime.fromisoformat(generated_utc).strftime("%Y-%m-%d %H:%M UTC")
    except ValueError:
        generated_label = generated_utc
    overview.cell(row + 5, 2, generated_label)
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 34

    note_row = row + 7
    overview.cell(note_row, 1, "How to read these results").font = Font(
        name="Arial", size=12, bold=True, color=NAVY
    )
    max_n = int(summary["n"].max()) if not summary.empty else 0
    interpretation = (
        "Lines show the median across repeated stochastic runs. Shaded ranges show the "
        "5th–95th percentiles of simulated trajectories; they are not confidence intervals."
        if max_n > 1 else
        "Each line is one simulated trajectory. A single run does not estimate stochastic variation."
    )
    overview.cell(note_row + 1, 1, interpretation)
    overview.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 2, end_column=8)
    overview.cell(note_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    overview.cell(note_row + 1, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    available = set(summary["metric"]) if not summary.empty else set()
    wanted = list(figure_metrics or (
        "new_infections", "new_diagnoses", "PLHIV", "on_art_stock"
    ))
    selected = [metric for metric in wanted if metric in available][:4]
    with tempfile.TemporaryDirectory(prefix="hivprep-workbook-") as temporary:
        for index, metric in enumerate(selected):
            png = Path(temporary) / f"{index}-{metric}.png"
            if not _chart_png(bundle, metric, png):
                continue
            # openpyxl only closes the temporary PIL handle automatically when
            # it receives a string path. Passing pathlib.Path leaks the handle
            # until garbage collection and fails the strict warning gate.
            image = Image(str(png))
            image.width = 680
            image.height = 340
            anchor = f"{'A' if index % 2 == 0 else 'J'}{note_row + 4 + (index // 2) * 19}"
            overview.add_image(image, anchor)

        annual = workbook.create_sheet("Annual results")
        _prepare_sheet(annual, freeze="A8")
        annual.sheet_properties.tabColor = "4C78A8"
        start = _title(
            annual, "Annual results",
            "Exact annual summaries by model setup. Percentiles describe stochastic runs.",
        )
        annual_rows = []
        ordered_summary = summary.copy()
        if not ordered_summary.empty:
            rank = {metric: index for index, metric in enumerate(KEY_METRICS)}
            ordered_summary["_rank"] = ordered_summary["metric"].map(
                lambda metric: rank.get(str(metric), len(rank))
            )
            ordered_summary = ordered_summary.sort_values(
                ["_rank", "metric", "route_id", "year"], kind="stable"
            )
        for record in ordered_summary.to_dict("records"):
            annual_rows.append([
                _route_label(record.get("route_id")), record.get("year"),
                CHART_LABELS.get(str(record.get("metric")), record.get("label")),
                _unit_label(record.get("metric"), record.get("unit")),
                record.get("n"), record.get("median"),
                record.get("q05"), record.get("q95"), record.get("minimum"),
                record.get("maximum"), AGGREGATION_LABELS.get(
                    str(record.get("annual_aggregation")), record.get("annual_aggregation")
                ),
            ])
        end = _write_table(
            annual, start,
            ["Model setup", "Year", "Outcome", "Unit", "Runs", "Median",
             "5th percentile", "95th percentile", "Minimum", "Maximum", "Annual calculation"],
            annual_rows, table_name="AnnualResults",
            widths=[20, 10, 34, 18, 10, 16, 16, 16, 16, 16, 20],
        )
        _format_number_cells(annual, start + 1, end, [6, 7, 8, 9, 10])

        individual = workbook.create_sheet("Individual runs")
        _prepare_sheet(individual, freeze="A8")
        start = _title(
            individual, "Individual runs",
            "Normalized annual values for every seed and model setup.",
        )
        ordered_individual = bundle.annual_long.copy()
        if not ordered_individual.empty:
            rank = {metric: index for index, metric in enumerate(KEY_METRICS)}
            ordered_individual["_rank"] = ordered_individual["metric"].map(
                lambda metric: rank.get(str(metric), len(rank))
            )
            ordered_individual = ordered_individual.sort_values(
                ["_rank", "metric", "run_id", "year"], kind="stable"
            )
        individual_rows = []
        for record in ordered_individual.to_dict("records"):
            individual_rows.append([
                record.get("run_id"), _route_label(record.get("route_id")),
                record.get("model"), record.get("p"), record.get("seed"),
                record.get("year"), CHART_LABELS.get(
                    str(record.get("metric")), record.get("label")
                ), record.get("value"),
                _unit_label(record.get("metric"), record.get("unit")),
                AGGREGATION_LABELS.get(
                    str(record.get("annual_aggregation")), record.get("annual_aggregation")
                ), record.get("source_file"),
            ])
        end = _write_table(
            individual, start,
            ["Run", "Model setup", "Model", "p", "Seed", "Year", "Outcome",
             "Value", "Unit", "Annual calculation", "Source file"],
            individual_rows, table_name="IndividualRuns",
            widths=[28, 18, 10, 10, 12, 10, 34, 16, 18, 20, 42],
        )
        _format_number_cells(individual, start + 1, end, [4, 8])

        configuration = workbook.create_sheet("Configuration")
        _prepare_sheet(configuration, freeze="A8")
        start = _title(
            configuration, "Configuration",
            "Parameters embedded in completed local runs.",
        )
        config_rows = _configuration_rows(bundle)
        if config_rows:
            _write_table(
                configuration, start,
                ["Group", "Parameter", "Starting value", "Run value", "Unit", "Change"],
                config_rows, table_name="ConfigurationValues",
                widths=[26, 48, 18, 18, 22, 14],
            )
        else:
            configuration.cell(start, 1, "No effective configuration was embedded in this result source.")
            configuration.cell(start, 1).fill = PatternFill("solid", fgColor=AMBER)

        provenance = workbook.create_sheet("Provenance")
        _prepare_sheet(provenance, freeze="A8")
        start = _title(
            provenance, "Provenance",
            "Exact identifiers and file hashes are retained here for reproducibility.",
        )
        run_rows = [[
            run.run_id, run.model, run.p, run.seed, run.status, run.classification,
            _display_path(run.manifest_path, bundle) if run.manifest_path else "",
        ] for run in bundle.runs]
        end = _write_table(
            provenance, start,
            ["Run", "Model", "p", "Seed", "Raw status", "Raw classification", "Manifest"],
            run_rows, table_name="RunProvenance",
            widths=[30, 10, 10, 12, 42, 36, 54],
        )
        source_start = end + 3
        provenance.cell(source_start - 1, 1, "Source files").font = Font(
            name="Arial", size=12, bold=True, color=NAVY
        )
        source_rows = [[
            table.relative_path, table.role, len(table.frame), len(table.frame.columns),
            table.size_bytes, table.sha256,
        ] for table in bundle.tables]
        _write_table(
            provenance, source_start,
            ["File", "Role", "Rows", "Columns", "Bytes", "SHA-256"],
            source_rows, table_name="SourceFiles",
            widths=[48, 30, 12, 12, 14, 68],
        )

        fit_path = bundle.source_path / "target_fit_candidate_1787.csv"
        if fit_path.is_file():
            fit = pd.read_csv(fit_path)
            calibration = workbook.create_sheet("Calibration fit")
            _prepare_sheet(calibration, freeze="A8")
            start = _title(
                calibration, "Calibration fit",
                "Primary threshold: I-max <= 3. Color repeats the text result and is not the only cue.",
            )
            fit_rows = []
            for record in fit.to_dict("records"):
                value = record.get("primary_implausibility")
                fit_rows.append([
                    record.get("target_key"), record.get("observed"),
                    record.get("simulated_mean"), record.get("signed_percentage_error"),
                    value, "Within threshold" if value is not None and float(value) <= 3 else "Above threshold",
                ])
            end = _write_table(
                calibration, start,
                ["Target", "Observed", "Model mean", "Difference (%)", "Implausibility", "Threshold"],
                fit_rows, table_name="CalibrationFit",
                widths=[42, 16, 16, 18, 18, 20],
            )
            _format_number_cells(calibration, start + 1, end, [2, 3, 4, 5])
            calibration.conditional_formatting.add(
                f"E{start + 1}:E{end}",
                CellIsRule(operator="lessThanOrEqual", formula=["3"], fill=PatternFill("solid", fgColor=GREEN)),
            )
            calibration.conditional_formatting.add(
                f"E{start + 1}:E{end}",
                CellIsRule(operator="greaterThan", formula=["3"], fill=PatternFill("solid", fgColor=RED)),
            )

        for ws in workbook.worksheets:
            ws.sheet_view.zoomScale = 90
            ws.sheet_properties.outlinePr.summaryBelow = True
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(output)
    return output
