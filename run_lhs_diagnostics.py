"""
Build post-run diagnostics for staged LHS calibration.

The diagnostics use the pipeline status JSON as the entry point, summarize the
ranked draw CSVs, rerun the final-stage top configs to recover annual model
outputs, and write an appendix-friendly Excel workbook plus a portable HTML
report.
"""

from __future__ import annotations

import argparse
import contextlib
import html
import io
import json
import math
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from calib_extractor import extract_annual_outputs
from calib_scorer import score_draw
from run_lhs_calibration import CALIB_END_YEAR, CALIB_START_YEAR
from run_v3 import SimulationEngineV3


PIPELINE_STAGES = ("backbone", "diagnosis", "art", "joint")
OUTPUT_ORDER = (
    "PLHIV",
    "new_infections",
    "aids_deaths",
    "new_diagnoses",
    "new_diagnoses_msm",
    "new_aids_cases",
    "on_art_stock",
)
OUTPUT_LABELS = {
    "PLHIV": "PLHIV",
    "new_infections": "New infections",
    "aids_deaths": "AIDS deaths",
    "new_diagnoses": "New diagnoses",
    "new_diagnoses_msm": "New diagnoses MSM",
    "new_aids_cases": "New AIDS cases",
    "on_art_stock": "On ART stock",
}
NAVY = "1F3864"
BLUE = "2F75B5"
ORANGE = "ED7D31"
GREY = "666666"
LIGHT_GREY = "F2F2F2"
GREEN = "C6EFCE"
AMBER = "FFEB9C"
RED = "FFC7CE"


def read_json(path: str | Path) -> dict:
    with Path(path).open(encoding="utf-8") as f:
        return json.load(f)


def write_json(path: str | Path, payload: dict) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2)
        f.write("\n")


def finite_series(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan).dropna()


def run_config_outputs(
    config_path: Path,
    seed: int,
    deterministic_mode: bool | None = None,
) -> dict[str, dict[int, float]]:
    config = read_json(config_path)
    config.setdefault("run_control", {}).setdefault("start_year", {})["value"] = CALIB_START_YEAR
    config.setdefault("run_control", {}).setdefault("end_year", {})["value"] = CALIB_END_YEAR + 1
    if deterministic_mode is not None:
        config.setdefault("run_control", {}).setdefault("deterministic_mode", {})["value"] = bool(deterministic_mode)
    config.setdefault("output", {}).setdefault("enabled", {})["value"] = False
    config.setdefault("output", {}).setdefault("terminal_summary", {})["value"] = False
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        result = SimulationEngineV3(config, seed=seed).run(run_number=1)
    return extract_annual_outputs(result, CALIB_START_YEAR, CALIB_END_YEAR)


def point_ll(model_value: float, estimate: float, sigma: float) -> float:
    sigma = max(float(sigma), 1e-9)
    resid = (float(model_value) - float(estimate)) / sigma
    return -0.5 * math.log(2.0 * math.pi * sigma * sigma) - 0.5 * resid * resid


def target_fit_frame(outputs: dict[str, dict[int, float]], targets: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for record in targets.to_dict("records"):
        model_output = record["model_output"]
        year = int(record["year"])
        estimate = float(record["estimate"])
        sigma = float(record["sigma"])
        weight = float(record.get("weight", 1.0))
        model_value = outputs.get(model_output, {}).get(year, np.nan)
        z_score = (model_value - estimate) / sigma if math.isfinite(model_value) else np.nan
        pct_diff = ((model_value - estimate) / estimate * 100.0) if estimate else np.nan
        rows.append(
            {
                "model_output": model_output,
                "year": year,
                "target_name": record.get("target_name", ""),
                "target_family": record.get("target_family", ""),
                "calibration_use": record.get("calibration_use", ""),
                "estimate": estimate,
                "model": model_value,
                "abs_diff": abs(model_value - estimate) if math.isfinite(model_value) else np.nan,
                "pct_diff": pct_diff,
                "sigma": sigma,
                "z_score": z_score,
                "abs_z": abs(z_score) if math.isfinite(z_score) else np.nan,
                "weight": weight,
                "weighted_ll": weight * point_ll(model_value, estimate, sigma)
                if math.isfinite(model_value)
                else np.nan,
            }
        )
    return pd.DataFrame(rows)


def outputs_frame(draw_outputs: dict[int, dict[str, dict[int, float]]]) -> pd.DataFrame:
    rows = []
    for rank, outputs in draw_outputs.items():
        for model_output, series in outputs.items():
            for year, value in series.items():
                rows.append(
                    {
                        "rank": rank,
                        "model_output": model_output,
                        "year": int(year),
                        "model": float(value),
                    }
                )
    return pd.DataFrame(rows)


def load_stage_tables(status: dict) -> tuple[dict[str, pd.DataFrame], dict[str, pd.DataFrame], dict[str, dict]]:
    ranked: dict[str, pd.DataFrame] = {}
    targets: dict[str, pd.DataFrame] = {}
    metadata: dict[str, dict] = {}
    for stage in PIPELINE_STAGES:
        stage_status = status.get("stages", {}).get(stage, {})
        if stage_status.get("status") != "completed":
            continue
        ranked_path = Path(stage_status["ranked_csv"])
        ranked[stage] = pd.read_csv(ranked_path)
        target_path = stage_status.get("normalized_targets")
        if target_path and Path(target_path).exists():
            targets[stage] = pd.read_csv(target_path)
        metadata_path = stage_status.get("metadata")
        if metadata_path and Path(metadata_path).exists():
            metadata[stage] = read_json(metadata_path)
    return ranked, targets, metadata


def stage_summary_frame(ranked: dict[str, pd.DataFrame], metadata: dict[str, dict]) -> pd.DataFrame:
    rows = []
    for stage, df in ranked.items():
        finite = finite_series(df["total_ll"])
        top20 = finite.head(20)
        target_summary = metadata.get(stage, {}).get("target_summary", [])
        rows.append(
            {
                "stage": stage,
                "n_draws": len(df),
                "finite_draws": len(finite),
                "collapsed_draws": int(df.get("collapsed", pd.Series(dtype=bool)).fillna(False).sum()),
                "best_draw": int(df.iloc[0]["draw_id"]),
                "best_total_ll": float(df.iloc[0]["total_ll"]),
                "median_total_ll": float(finite.median()) if len(finite) else np.nan,
                "p95_total_ll": float(finite.quantile(0.95)) if len(finite) else np.nan,
                "p99_total_ll": float(finite.quantile(0.99)) if len(finite) else np.nan,
                "top20_gap": float(top20.iloc[0] - top20.iloc[-1]) if len(top20) >= 20 else np.nan,
                "target_points": int(sum(item.get("n_targets", 0) for item in target_summary))
                if target_summary
                else np.nan,
            }
        )
    return pd.DataFrame(rows)


def component_frame(ranked: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows = []
    for stage, df in ranked.items():
        best = df.iloc[0]
        for col in [c for c in df.columns if c.startswith("ll_")]:
            rows.append(
                {
                    "stage": stage,
                    "component": col.removeprefix("ll_"),
                    "weighted_ll": float(best[col]),
                    "share_abs_total": abs(float(best[col])) / abs(float(best["total_ll"]))
                    if float(best["total_ll"])
                    else np.nan,
                }
            )
    return pd.DataFrame(rows)


def parameter_frame(ranked: dict[str, pd.DataFrame], metadata: dict[str, dict]) -> pd.DataFrame:
    rows = []
    for stage, df in ranked.items():
        space_path = metadata.get(stage, {}).get("calib_space_path")
        ranges: dict[str, tuple[float, float]] = {}
        if space_path and Path(space_path).exists():
            space = read_json(space_path)
            ranges = {
                p["name"]: (float(p["min"]), float(p["max"]))
                for p in space.get("parameters", [])
                if stage in p.get("stages", [])
            }
        best = df.iloc[0]
        for param, (lower, upper) in ranges.items():
            if param not in df.columns:
                continue
            value = float(best[param])
            pct = (value - lower) / (upper - lower) if upper > lower else np.nan
            rows.append(
                {
                    "stage": stage,
                    "parameter": param,
                    "value": value,
                    "min": lower,
                    "max": upper,
                    "range_pct": pct,
                    "edge_flag": "LOW" if pct <= 0.05 else "HIGH" if pct >= 0.95 else "",
                    "top20_mean": float(df.head(20)[param].mean()),
                    "top20_sd": float(df.head(20)[param].std()),
                }
            )
        if "activity_low" in df.columns:
            rows.append(
                {
                    "stage": stage,
                    "parameter": "activity_low",
                    "value": float(best["activity_low"]),
                    "min": np.nan,
                    "max": np.nan,
                    "range_pct": np.nan,
                    "edge_flag": "",
                    "top20_mean": float(df.head(20)["activity_low"].mean()),
                    "top20_sd": float(df.head(20)["activity_low"].std()),
                }
            )
    return pd.DataFrame(rows)


def residual_summary_frame(fit: pd.DataFrame) -> pd.DataFrame:
    return (
        fit.groupby("model_output")
        .agg(
            n=("year", "count"),
            mean_pct=("pct_diff", "mean"),
            median_abs_pct=("pct_diff", lambda s: s.abs().median()),
            max_abs_pct=("pct_diff", lambda s: s.abs().max()),
            mean_z=("z_score", "mean"),
            rms_z=("z_score", lambda s: math.sqrt(float(np.nanmean(np.square(s))))),
            weighted_ll=("weighted_ll", "sum"),
        )
        .reset_index()
        .sort_values("weighted_ll")
    )


def autosize(ws, max_width: int = 42) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = 10
        for cell in column:
            if cell.value is None:
                continue
            width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def style_header(ws) -> None:
    header_fill = PatternFill("solid", fgColor=NAVY)
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_df(ws, df: pd.DataFrame, start_row: int = 1, start_col: int = 1) -> None:
    for col_idx, col in enumerate(df.columns, start=start_col):
        ws.cell(start_row, col_idx, col)
    for row_idx, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col_idx, value in enumerate(row, start=start_col):
            if isinstance(value, (np.floating, np.integer)):
                value = float(value)
            if isinstance(value, float) and not math.isfinite(value):
                value = None
            ws.cell(row_idx, col_idx, value)


def add_table_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> Any:
    ws = wb.create_sheet(title)
    write_df(ws, df)
    style_header(ws)
    autosize(ws)
    thin = Side(style="thin", color="BFBFBF")
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    return ws


def add_summary_sheet(wb: Workbook, status: dict, stage_summary: pd.DataFrame, residual_summary: pd.DataFrame) -> None:
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "LHS Calibration Diagnostics"
    ws["A1"].font = Font(name="Arial", bold=True, size=16, color=NAVY)
    ws["A3"] = "Created"
    ws["B3"] = datetime.now().isoformat(timespec="seconds")
    ws["A4"] = "Pipeline status"
    ws["B4"] = status.get("status")
    ws["A5"] = "Initial config"
    ws["B5"] = status.get("initial_config")
    ws["A6"] = "Final config"
    ws["B6"] = status.get("final_config")

    ws["A8"] = "Stage score summary"
    ws["A8"].font = Font(name="Arial", bold=True, size=12, color=NAVY)
    write_df(ws, stage_summary, start_row=9)
    for cell in ws[9]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")

    start = 9 + len(stage_summary) + 3
    ws.cell(start, 1, "Final best residual summary").font = Font(name="Arial", bold=True, size=12, color=NAVY)
    write_df(ws, residual_summary, start_row=start + 1)
    for cell in ws[start + 1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
    autosize(ws, max_width=70)


def add_residual_heatmap(wb: Workbook, fit: pd.DataFrame) -> None:
    pivot = (
        fit.pivot_table(index="model_output", columns="year", values="z_score", aggfunc="mean")
        .reindex([o for o in OUTPUT_ORDER if o in set(fit["model_output"])])
        .reset_index()
    )
    ws = add_table_sheet(wb, "Residual_Heatmap", pivot)
    if ws.max_row > 1 and ws.max_column > 2:
        start = ws.cell(2, 2).coordinate
        end = ws.cell(ws.max_row, ws.max_column).coordinate
        ws.conditional_formatting.add(
            f"{start}:{end}",
            ColorScaleRule(
                start_type="num",
                start_value=-4,
                start_color="5B9BD5",
                mid_type="num",
                mid_value=0,
                mid_color="FFFFFF",
                end_type="num",
                end_value=4,
                end_color="F8696B",
            ),
        )
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell.number_format = "0.0"


def add_excel_charts(wb: Workbook, fit: pd.DataFrame, model_outputs: pd.DataFrame, component: pd.DataFrame) -> None:
    chart_ws = wb.create_sheet("Charts")
    chart_ws["A1"] = "Trajectory charts: target vs final best-fit model"
    chart_ws["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY)
    row_cursor = 3
    for output in [o for o in OUTPUT_ORDER if o in set(fit["model_output"])]:
        target = fit[fit["model_output"] == output][["year", "estimate", "model"]].sort_values("year")
        write_df(chart_ws, target, start_row=row_cursor, start_col=1)
        for cell in chart_ws[row_cursor]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        chart = LineChart()
        chart.title = OUTPUT_LABELS.get(output, output)
        chart.y_axis.title = "Count"
        chart.x_axis.title = "Year"
        data = Reference(chart_ws, min_col=2, max_col=3, min_row=row_cursor, max_row=row_cursor + len(target))
        cats = Reference(chart_ws, min_col=1, min_row=row_cursor + 1, max_row=row_cursor + len(target))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 7
        chart.width = 14
        chart_ws.add_chart(chart, f"E{row_cursor}")
        row_cursor += max(len(target) + 3, 18)

    bar_start = row_cursor + 1
    final_component = component[component["stage"] == "joint"].copy()
    write_df(chart_ws, final_component[["component", "weighted_ll"]], start_row=bar_start, start_col=1)
    bar = BarChart()
    bar.title = "Joint best likelihood contribution"
    bar.y_axis.title = "Weighted log-likelihood"
    data = Reference(chart_ws, min_col=2, min_row=bar_start, max_row=bar_start + len(final_component))
    cats = Reference(chart_ws, min_col=1, min_row=bar_start + 1, max_row=bar_start + len(final_component))
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.height = 8
    bar.width = 16
    chart_ws.add_chart(bar, f"E{bar_start}")
    autosize(chart_ws)


def build_excel_report(
    output_path: Path,
    status: dict,
    ranked: dict[str, pd.DataFrame],
    stage_summary: pd.DataFrame,
    component: pd.DataFrame,
    params: pd.DataFrame,
    fit: pd.DataFrame,
    residual_summary: pd.DataFrame,
    model_outputs: pd.DataFrame,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    add_summary_sheet(wb, status, stage_summary, residual_summary)
    add_table_sheet(wb, "Stage_Scores", stage_summary)
    add_table_sheet(wb, "LL_Contributions", component)
    add_table_sheet(wb, "Best_Params", params)
    add_table_sheet(wb, "Residual_Summary", residual_summary)
    add_table_sheet(wb, "Target_Fit", fit.sort_values(["model_output", "year"]))
    add_residual_heatmap(wb, fit)
    add_table_sheet(wb, "Top_Draws_Joint", ranked["joint"].head(100))
    add_table_sheet(wb, "Model_Outputs", model_outputs.sort_values(["rank", "model_output", "year"]))
    add_excel_charts(wb, fit, model_outputs, component)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def svg_line_chart(data: pd.DataFrame, title: str, width: int = 820, height: int = 260) -> str:
    if data.empty:
        return ""
    margin = {"left": 58, "right": 18, "top": 32, "bottom": 36}
    xs = data["year"].astype(float)
    ys = pd.concat([data["estimate"], data["model"]]).astype(float)
    x_min, x_max = float(xs.min()), float(xs.max())
    y_min, y_max = 0.0, float(ys.max()) * 1.08 if float(ys.max()) > 0 else 1.0

    def sx(x: float) -> float:
        return margin["left"] + (x - x_min) / max(x_max - x_min, 1) * (width - margin["left"] - margin["right"])

    def sy(y: float) -> float:
        return height - margin["bottom"] - (y - y_min) / max(y_max - y_min, 1) * (height - margin["top"] - margin["bottom"])

    def path_for(col: str) -> str:
        points = [f"{sx(float(row.year)):.1f},{sy(float(getattr(row, col))):.1f}" for row in data.itertuples()]
        return "M " + " L ".join(points)

    target_path = path_for("estimate")
    model_path = path_for("model")
    y_ticks = [0, y_max / 2, y_max]
    x_ticks = [x_min, (x_min + x_max) / 2, x_max]
    tick_svg = []
    for y in y_ticks:
        tick_svg.append(f'<line x1="{margin["left"]}" y1="{sy(y):.1f}" x2="{width-margin["right"]}" y2="{sy(y):.1f}" stroke="#e6e6e6"/>')
        tick_svg.append(f'<text x="{margin["left"]-8}" y="{sy(y)+4:.1f}" text-anchor="end" class="tick">{y:,.0f}</text>')
    for x in x_ticks:
        tick_svg.append(f'<text x="{sx(x):.1f}" y="{height-10}" text-anchor="middle" class="tick">{int(round(x))}</text>')
    return f"""
<svg viewBox="0 0 {width} {height}" class="chart" role="img" aria-label="{html.escape(title)}">
  <text x="0" y="18" class="chart-title">{html.escape(title)}</text>
  {''.join(tick_svg)}
  <line x1="{margin["left"]}" y1="{height-margin["bottom"]}" x2="{width-margin["right"]}" y2="{height-margin["bottom"]}" stroke="#999"/>
  <line x1="{margin["left"]}" y1="{margin["top"]}" x2="{margin["left"]}" y2="{height-margin["bottom"]}" stroke="#999"/>
  <path d="{target_path}" fill="none" stroke="#{ORANGE}" stroke-width="2.2"/>
  <path d="{model_path}" fill="none" stroke="#{BLUE}" stroke-width="2.2"/>
  <text x="{width-170}" y="24" class="legend" fill="#{ORANGE}">Target</text>
  <text x="{width-90}" y="24" class="legend" fill="#{BLUE}">Model</text>
</svg>"""


def residual_color(z: float) -> str:
    if not math.isfinite(z):
        return "#eeeeee"
    z = max(-4.0, min(4.0, z))
    if z >= 0:
        t = z / 4.0
        r, g, b = 255, int(255 - 120 * t), int(255 - 120 * t)
    else:
        t = abs(z) / 4.0
        r, g, b = int(255 - 120 * t), int(255 - 70 * t), 255
    return f"#{r:02x}{g:02x}{b:02x}"


def heatmap_html(fit: pd.DataFrame) -> str:
    pivot = fit.pivot_table(index="model_output", columns="year", values="z_score", aggfunc="mean")
    years = list(pivot.columns)
    rows = ['<table class="heatmap"><thead><tr><th>Output</th>']
    rows.extend(f"<th>{int(year)}</th>" for year in years)
    rows.append("</tr></thead><tbody>")
    for output in [o for o in OUTPUT_ORDER if o in pivot.index]:
        rows.append(f"<tr><th>{html.escape(OUTPUT_LABELS.get(output, output))}</th>")
        for year in years:
            value = pivot.loc[output, year]
            label = "" if pd.isna(value) else f"{value:.1f}"
            rows.append(f'<td style="background:{residual_color(float(value) if pd.notna(value) else float("nan"))}">{label}</td>')
        rows.append("</tr>")
    rows.append("</tbody></table>")
    return "".join(rows)


def build_html_report(
    output_path: Path,
    status: dict,
    stage_summary: pd.DataFrame,
    component: pd.DataFrame,
    params: pd.DataFrame,
    fit: pd.DataFrame,
    residual_summary: pd.DataFrame,
) -> None:
    charts = []
    for output in [o for o in OUTPUT_ORDER if o in set(fit["model_output"])]:
        chart_data = fit[fit["model_output"] == output][["year", "estimate", "model"]].sort_values("year")
        charts.append(svg_line_chart(chart_data, OUTPUT_LABELS.get(output, output)))
    worst = fit.sort_values("abs_z", ascending=False).head(20)
    edge_params = params[params["edge_flag"] != ""].copy()
    if edge_params.empty:
        edge_html = "<p>No best-fit parameters are within 5% of a configured bound.</p>"
    else:
        edge_html = edge_params[["stage", "parameter", "value", "min", "max", "range_pct", "edge_flag"]].to_html(
            index=False, classes="data", float_format=lambda x: f"{x:.3g}", escape=True
        )
    html_text = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>LHS Calibration Diagnostics</title>
<style>
body {{ font-family: Arial, sans-serif; margin: 28px; color: #1f1f1f; }}
h1, h2 {{ color: #{NAVY}; }}
.meta {{ color: #{GREY}; }}
.grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(360px, 1fr)); gap: 18px; }}
.panel {{ border: 1px solid #d9d9d9; border-radius: 6px; padding: 14px; background: #fff; }}
table.data {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
table.data th {{ background: #{NAVY}; color: #fff; text-align: left; }}
table.data th, table.data td {{ border: 1px solid #d0d0d0; padding: 6px 8px; }}
.heatmap {{ border-collapse: collapse; font-size: 11px; width: 100%; }}
.heatmap th, .heatmap td {{ border: 1px solid #d0d0d0; padding: 4px; text-align: center; }}
.heatmap th {{ background: #{NAVY}; color: white; position: sticky; left: 0; }}
.chart {{ width: 100%; height: auto; }}
.chart-title {{ font-weight: 700; fill: #{NAVY}; font-size: 14px; }}
.tick {{ fill: #555; font-size: 10px; }}
.legend {{ font-size: 12px; font-weight: 700; }}
</style>
</head>
<body>
<h1>LHS Calibration Diagnostics</h1>
<p class="meta">Created {datetime.now().isoformat(timespec="seconds")} from {html.escape(str(status.get("output_dir", "")))}</p>
<p><strong>Final config:</strong> {html.escape(str(status.get("final_config", "")))}</p>
<h2>Stage Scores</h2>
{stage_summary.to_html(index=False, classes="data", float_format=lambda x: f"{x:.2f}", escape=True)}
<h2>Final Residual Summary</h2>
{residual_summary.to_html(index=False, classes="data", float_format=lambda x: f"{x:.2f}", escape=True)}
<h2>Trajectory Fit</h2>
<div class="grid">{''.join(f'<div class="panel">{chart}</div>' for chart in charts)}</div>
<h2>Residual Heatmap</h2>
{heatmap_html(fit)}
<h2>Likelihood Contributions</h2>
{component[component["stage"] == "joint"].to_html(index=False, classes="data", float_format=lambda x: f"{x:.2f}", escape=True)}
<h2>Parameters Near Bounds</h2>
{edge_html}
<h2>Worst Target-Year Residuals</h2>
{worst[["model_output", "year", "estimate", "model", "pct_diff", "z_score", "weighted_ll", "target_name"]].to_html(index=False, classes="data", float_format=lambda x: f"{x:.2f}", escape=True)}
</body>
</html>
"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(html_text, encoding="utf-8")


def build_diagnostics(args: argparse.Namespace) -> dict[str, str]:
    status_path = Path(args.status).resolve()
    status = read_json(status_path)
    stages = status.get("stages", {})
    stages_completed = all(
        stages.get(stage, {}).get("status") == "completed"
        for stage in PIPELINE_STAGES
    )
    if status.get("status") != "completed" and not stages_completed:
        raise RuntimeError(f"Pipeline is not completed: {status_path}")
    ranked, targets, metadata = load_stage_tables(status)
    if "joint" not in ranked or "joint" not in targets:
        raise RuntimeError("Diagnostics require completed joint ranked results and normalized targets.")

    joint = ranked["joint"]
    top_n = min(int(args.top_n), len(joint))
    final_stage_status = status["stages"]["joint"]
    stage_seed = int(final_stage_status["seed"])
    best_config_paths = [
        Path(path)
        for path in final_stage_status.get("best_fit_configs", [])
        if path and Path(path).exists()
    ]
    if not best_config_paths:
        stage_dir = Path(final_stage_status["output_dir"])
        best_config_paths = [
            stage_dir / f"lhs_best_fit_rank_{rank:03d}.json"
            for rank in range(1, top_n + 1)
        ]
    joint_metadata = metadata.get("joint", {})
    deterministic_mode = joint_metadata.get("deterministic_calibration")
    family_balanced = bool(joint_metadata.get("family_balanced_score", False))
    reranked = int(joint_metadata.get("rerank_replicates", 1) or 1) > 1
    draw_outputs: dict[int, dict[str, dict[int, float]]] = {}
    for rank in range(1, top_n + 1):
        row = joint.iloc[rank - 1]
        if rank > len(best_config_paths):
            continue
        config_path = best_config_paths[rank - 1]
        if not config_path.exists():
            continue
        draw_outputs[rank] = run_config_outputs(
            config_path,
            seed=stage_seed + int(row["draw_id"]),
            deterministic_mode=deterministic_mode,
        )

    if 1 not in draw_outputs:
        raise RuntimeError("Diagnostics could not locate the joint rank-001 best-fit config.")

    best_outputs = draw_outputs[1]
    target_records = targets["joint"]
    score = score_draw(
        best_outputs,
        target_records.to_dict("records"),
        balance_families=family_balanced,
    )
    if not reranked and abs(float(score["total_ll"]) - float(joint.iloc[0]["total_ll"])) > 1e-6:
        raise RuntimeError(
            "Rerun score does not match ranked CSV for joint best draw: "
            f"{score['total_ll']} vs {joint.iloc[0]['total_ll']}"
        )

    fit = target_fit_frame(best_outputs, target_records)
    stage_summary = stage_summary_frame(ranked, metadata)
    component = component_frame(ranked)
    params = parameter_frame(ranked, metadata)
    residual_summary = residual_summary_frame(fit)
    model_outputs = outputs_frame(draw_outputs)

    output_dir = Path(args.output_dir).resolve() if args.output_dir else status_path.parent / "diagnostics"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_path = output_dir / f"lhs_diagnostics_{stamp}.xlsx"
    html_path = output_dir / f"lhs_diagnostics_{stamp}.html"
    build_excel_report(excel_path, status, ranked, stage_summary, component, params, fit, residual_summary, model_outputs)
    build_html_report(html_path, status, stage_summary, component, params, fit, residual_summary)
    manifest_path = output_dir / f"lhs_diagnostics_manifest_{stamp}.json"
    manifest = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "status_path": str(status_path),
        "top_n_rerun": top_n,
        "excel_report": str(excel_path),
        "html_report": str(html_path),
    }
    write_json(manifest_path, manifest)
    return manifest


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build Excel and HTML diagnostics for a completed LHS pipeline.")
    parser.add_argument(
        "--status",
        default=str(ROOT / "outputs" / "calibration" / "lhs_pipeline" / "lhs_pipeline_status.json"),
        help="Path to lhs_pipeline_status.json.",
    )
    parser.add_argument("--output-dir", default=None, help="Directory for diagnostics reports.")
    parser.add_argument("--top-n", type=int, default=5, help="Number of final-stage top configs to rerun.")
    args = parser.parse_args(argv)
    if args.top_n <= 0:
        parser.error("--top-n must be positive")
    return args


def main(argv: Sequence[str] | None = None) -> int:
    try:
        manifest = build_diagnostics(parse_args(argv))
    except BaseException as exc:
        print(f"Diagnostics failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    print(f"Excel diagnostics: {manifest['excel_report']}")
    print(f"HTML diagnostics: {manifest['html_report']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
